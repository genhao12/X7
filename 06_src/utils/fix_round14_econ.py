"""Round14 经济数值更新（degradation_summary/cost_sensitivity 已重算完成）。
更新 P92 经济段、P88/P90 成本口径、汇总 Table11/Table7、台账 S10。"""
import docx
from pathlib import Path

ROOT = Path(r"d:\matlab\C")

def set_para(p, new):
    for r in p.runs: r.text = ""
    if p.runs: p.runs[0].text = new
    else: p.add_run(new)

# ───────── 主报告 ─────────
p = ROOT / "09_docs" / "工作区全面总结报告.docx"
d = docx.Document(str(p))

# 1) P88 成本口径注（更新：电力电子已排除 H2 罐）
for para in d.paragraphs:
    if "本表单价为设备本体（components）CAPEX" in para.text:
        set_para(para, "注：本表单价为设备本体（components）CAPEX，不含安装（round14 统一口径）；"
                        "电力电子（逆变器/整流器/DC-DC）按 PV+BAT+PEMFC+PEMEL 子小计的 15% 另计（不含 H₂ 储罐），"
                        "安装按全部设备 subtotal 的 20% 另计（npc_lite.compute_equipment_capex，round14 已修复）。")
        print("P88 成本口径注更新")
        break

# 2) P90 成本构成口径（更新：H2 罐不再被加 15%）
for para in d.paragraphs:
    if "成本构成口径（compute_equipment_capex）" in para.text:
        set_para(para, "成本构成口径（compute_equipment_capex，round14 修复）：设备单价为设备本体 CAPEX"
                        "（PV 含组件+支架、BAT 含电芯+BMS+Pack、FC/EL 含电堆+BOP、H₂ 含储罐，均不含安装）；"
                        "电力电子（逆变器/整流器/DC-DC）按 PV+BAT+PEMFC+PEMEL 子小计的 15% 另计"
                        "（power_electronics_fraction=0.15，**不再施加于 H₂ 储罐**，消除储罐高估）；"
                        "安装按全部设备 subtotal 的 20% 另计（installation_fraction=0.20，设备单价不含安装，无重复计费）；"
                        "FC/EL 的 BOP 已含在自身单价中（电堆+BOP），不重复计入 15%；"
                        "运维按各自 CAPEX 比例估算（非实际 O&M 合同）；无税费、无保险、无土地成本；"
                        "残值设为 0（BAT 梯次利用、H₂ 罐 30 年残值均未计）。")
        print("P90 成本构成口径更新")
        break

# 3) P92 经济段
for para in d.paragraphs:
    if "基线 NPC 31.5~39.3 MCNY" in para.text:
        set_para(para, "经济性基于 EFC 循环损伤模型重算（degradation_summary.json economics，round14 成本口径修复后重跑）："
                        "5 配置 × 16 情景 = 80 组（economics_allowed=True，门控已通过）。"
                        "基线 NPC 31.0~38.6 MCNY、LCOE 5.46~6.78 CNY/kWh（compact 最低 5.4593、"
                        "hydrogen_dominant 最高 6.7835，见 Table14；round14 修复后较旧值 5.55~6.90 全面下降约 0.1~0.15）；"
                        "单设备 +30% 敏感性：BAT 为 3/5 配置第一敏感项（compact/battery_dominant/critical，LCOE +9.4~11.7%）、"
                        "PEMEL 为 2/5（balanced/hydrogen_dominant，+10.0~12.0%）；"
                        "BOP 剥离降 LCOE 17.0~18.3%（较 round12 旧值 19.9~21.2% 下降，因电力电子基数缩小）；"
                        "全设备同步缩放 NPC 严格线性。经济结论限定为情景评价（成本参数边界未补全，见下）。")
        print("P92 经济段更新")
        break

d.save(str(p))
print("=== 主报告保存 ===")

# ───────── 汇总报告 ─────────
p = ROOT / "09_docs" / "工作区运行结果汇总.docx"
d = docx.Document(str(p))

# 4) 汇总 Table11 成本敏感性主要结果
for tb in d.tables:
    if len(tb.rows) > 1 and "run_cost_sensitivity" in tb.rows[1].cells[1].text:
        for row in tb.rows:
            if row.cells[0].text.strip() == "主要结果":
                if "5.55~6.90" in row.cells[1].text:
                    row.cells[1].text = ("基线 LCOE 5.46~6.78 CNY/kWh（round14 成本口径修复后重跑：compact 5.4593、"
                                         "hydrogen_dominant 6.7835）；单设备+30%：BAT 3/5 配置第一敏感项（+9.4~11.7%）、"
                                         "PEMEL 2/5（+10.0~12.0%）；BOP 剥离降 LCOE 17.0~18.3%；全设备包络线性")
                    print("汇总 Table11 主要结果更新")
                break
        break

# 5) 汇总 Table7 M4 经济中的 Stage3 EFC 引用
for tb in d.tables:
    if len(tb.rows) > 1 and "npc_lite" in tb.rows[1].cells[1].text:
        for row in tb.rows:
            if row.cells[0].text.strip() == "主要结果" and "Stage2c 5 配置 EFC 经济" in row.cells[1].text:
                row.cells[1].text = row.cells[1].text.replace(
                    "Stage2c 5 配置 EFC 经济：NPC 31.5~39.3M、LCOE 5.55~6.90",
                    "Stage2c 5 配置 EFC 经济（round14）：NPC 31.0~38.6M、LCOE 5.46~6.78")
                print("汇总 Table7 EFC 经济引用更新")
        break

d.save(str(p))
print("=== 汇总报告保存 ===")

# ───────── 台账 ─────────
p = ROOT / "09_docs" / "统一结果台账.docx"
d = docx.Document(str(p))

# 6) 台账 S10 成本敏感性基线
for row in d.tables[1].rows:
    if row.cells[0].text.strip().startswith("S10"):
        for ci, c in enumerate(row.cells):
            if "5.55" in c.text or "6.90" in c.text:
                c.text = c.text.replace("5.55~6.90", "5.46~6.78").replace("5.55–6.90", "5.46~6.78")
                print(f"台账 S10 C{ci} 基线更新")
        # 补 round14 说明
        if "round14" not in row.cells[7].text:
            row.cells[7].text = row.cells[7].text.rstrip() + "；round14 成本口径修复（电力电子排除 H₂ 罐、设备单价去含安装）后重跑，基线 LCOE 5.46~6.78"
            print("台账 S10 备注补 round14")
        break

d.save(str(p))
print("=== 台账保存 ===")
print("Round14 经济更新完成")
