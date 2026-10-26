"""Rebuild the workspace summary DOCX with verified Stage 0-2 evidence."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / "09_docs" / "工作区全面总结报告.docx"
RESULT = ROOT / "07_results" / "rectification_v2"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_GRAY = "F2F4F7"
CAUTION = "FFF2CC"
RISK = "FCE4D6"
INK = "1F1F1F"
MUTED = "666666"


def set_east_asia(run, name="等线"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths, font_size=8.5, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(text)
        set_cell_shading(cell, header_fill)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_east_asia(run)
                run.bold = True
                run.font.size = Pt(font_size)
    for row_values in rows:
        row = table.add_row()
        for index, text in enumerate(row_values):
            cell = row.cells[index]
            cell.text = str(text)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_east_asia(run)
                    run.font.size = Pt(font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_border_bottom(paragraph, color=BLUE, size=12):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_page_number(paragraph):
    run = paragraph.add_run("第 ")
    set_east_asia(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    tail = paragraph.add_run(" 页")
    set_east_asia(tail)


def add_paragraph(doc, text, bold_prefix=None, color=None, keep=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        set_east_asia(first)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_east_asia(rest)
    else:
        run = paragraph.add_run(text)
        set_east_asia(run)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_east_asia(run)
    return paragraph


def new_page(doc):
    doc.add_page_break()


def load_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def build():
    audit = load_json(RESULT / "input_contract" / "input_audit.json")
    gate = load_json(RESULT / "capacity_gate_summary.json")
    history = load_json(RESULT / "historical_result_status.json")
    s12 = load_json(RESULT / "tests" / "s12_regression.json")
    best_id = "extended_pv__pv320__bat180__fc45__el45__h210"
    best = load_json(RESULT / "extended_domain" / "checkpoints" / f"{best_id}.json")

    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for header_part in (section.header, section.even_page_header):
        header = header_part.paragraphs[0]
        header.text = "连续年度可行基线尚未建立  |  整改阶段结果与当前阻断报告"
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in header.runs:
            set_east_asia(run)
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(MUTED)
    for footer_part in (section.footer, section.even_page_footer):
        footer = footer_part.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_page_number(footer)
        for run in footer.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("离网光伏–氢–电池混合储能微电网")
    set_east_asia(run)
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("工作区全面总结报告｜阶段0–2正式实施版")
    set_east_asia(run)
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    for label, value in (
        ("文档定位", "整改阶段结果与当前阻断报告"),
        ("更新日期", "2026-08-02"),
        ("覆盖范围", "代码、配置、测试、结果目录与DOCX同步"),
        ("当前总判定", "连续年度可行基线尚未建立"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        r1 = paragraph.add_run(f"{label}：")
        set_east_asia(r1)
        r1.bold = True
        r2 = paragraph.add_run(value)
        set_east_asia(r2)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(14)
    set_border_bottom(rule)
    verdict = add_paragraph(
        doc,
        "阶段2容量准入结果：主研究域0组进入周期精筛；扩展域640组进入周期精筛、0组准入；旧C5两组诊断配置均不准入。退化、六路径、Shapley、L4、经济与双层不确定性实算继续冻结。",
        keep=True,
    )
    verdict.paragraph_format.space_before = Pt(10)
    verdict.paragraph_format.space_after = Pt(12)
    for run in verdict.runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("9B1C1C")

    heading(doc, "一、执行结论与依赖门控", 1)
    add_table(
        doc,
        ["对象", "执行结果", "论文状态"],
        [
            ["PV输入契约", "233 kW参考容量；候选容量按比例缩放并限幅", "阶段1通过"],
            ["主研究域", "90–150 kWp，2240组；全部未通过必要能量筛选", "原设计域不可行"],
            ["扩展容量情景", "160–320 kWp，2880组；640组精筛，0组准入", "容量跃迁尚未形成可行解"],
            ["旧C5诊断", "300/800/100/100/500或1000；2组均不准入", "仅历史对照"],
            ["三年周期稳态", "仅要求2016–2018三年块首尾SOC/H₂闭合", "单年漂移仅诊断"],
            ["退化/路径/Shapley/L4/经济", "未获容量准入", "冻结"],
            ["双层不确定性", "样本量与停止规则已配置，未执行", "冻结"],
        ],
        [1800, 4560, 3000],
    )

    new_page(doc)
    heading(doc, "二、阶段1：PV数据契约与输入审计", 1)
    add_paragraph(
        doc,
        "PV参考装机容量 = 233 kW。候选可用功率按 P_PV(t)=min[P_nom, max(0,P_raw(t))×P_nom/233] 计算；逆变器额定功率等于候选P_nom。当前原始曲线未检测到连续高功率平台，因此允许比例缩放；一旦检测到限发或削顶平台，P_nom>233 kWp的放大立即中止，并切换至辐照度/容量因子重建。负荷序列保持原值，不缩放、不改符号。",
    )
    add_table(
        doc,
        ["年份", "点数", "原始PV最大值(kW)", "负荷(MWh)", "150 kWp PV(MWh)", "PV/负荷"],
        [
            [year, row["points"], f'{row["pv_raw_max_kw"]:.3f}', f'{row["load_energy_kwh"]/1000:.3f}', f'{row["pv_150_energy_kwh"]/1000:.3f}', f'{100*row["pv_150_to_load_ratio"]:.3f}%']
            for year, row in audit["years"].items()
        ],
        [900, 900, 1740, 1680, 2100, 2040],
    )
    add_paragraph(
        doc,
        f"输入文件SHA-256：{audit['input_hash']}。2016为闰年35136点，2017/2018各35040点，均为完整15 min本地民用时序；原始数据属于合成重建评价序列，容量准入结果可用于方法与边界诊断，禁止替代真实站点多年度统计。",
        color=MUTED,
    )

    heading(doc, "2.1 研究域严格隔离", 2)
    add_table(
        doc,
        ["域", "PV(kWp)", "BAT(kWh)", "FC(kW)", "EL(kW)", "H₂(kg)", "组合数"],
        [
            ["主研究域", "90:10:150", "90/120/150/180", "30/35/40/45", "25/30/35/40/45", "10/20/30/40", "2240"],
            ["扩展容量情景", "160:20:320", "90/120/150/180", "30/35/40/45", "25/30/35/40/45", "10/20/30/40", "2880"],
            ["旧C5诊断", "300", "800", "100", "100", "500/1000", "2"],
        ],
        [1440, 1200, 1560, 1320, 1440, 1320, 1080],
        font_size=7.8,
    )
    add_paragraph(doc, "C5=300 kWp仅位于旧C5诊断/扩展语境，不与90–150 kWp主样本混合，也不替代冻结设计。")

    new_page(doc)
    heading(doc, "三、阶段2：容量搜索、周期稳态与准入", 1)
    add_paragraph(
        doc,
        "执行链固定为：必要能量筛选 → 2017单年粗排 → 2016–2018三年块周期精筛 → 最多8进程并行。粗排遍历SOC={0.20,0.60,0.95}与H₂初始比例={0.20,0.60,1.00}的9个组合，只决定精筛初态，不直接淘汰候选。周期精筛连续传递年际状态并最多重复50轮；收敛阈值为三年块首尾SOC绝对差与H₂额定比例差均≤0.001。单个自然年的库存变化保留为诊断字段，不是准入条件。",
    )
    add_table(
        doc,
        ["域", "总组合", "必要能量通过", "周期精筛", "准入", "下游状态"],
        [
            ["主研究域", "2240", "0", "0", "0", "冻结"],
            ["扩展容量情景", "2880", "640", "640", "0", "冻结"],
            ["旧C5诊断", "2", "2", "2", "0", "仅对照"],
        ],
        [1800, 1320, 1800, 1440, 1200, 1800],
    )
    add_paragraph(
        doc,
        "主研究域的150 kWp上边界在2016/2017/2018的PV/负荷比分别为51.134%、54.687%、59.625%，低于99%必要能量阈值；这是安全淘汰条件，未将其误作充分可行性条件。",
    )

    heading(doc, "3.1 扩展域最优诊断配置仍不准入", 2)
    add_paragraph(
        doc,
        "扩展域三年LPSP最低配置为PV 320 kWp、BAT 180 kWh、FC 45 kW、EL 45 kW、H₂ 10 kg。该配置2轮达到三年块首尾SOC/H₂差均为0，但三年LPSP=40.2684%；闭合发生在SOC=0.15、H₂=1.0 kg下限状态，说明周期闭合不等于可靠供电。",
    )
    yearly = best["metrics"]["yearly"]
    add_table(
        doc,
        ["年份", "LPSP", "关键负荷满足率", "块内SOC变化", "块内H₂变化", "H₂下限时长(h)"],
        [
            [year, f'{100*row["lpsp"]:.4f}%', f'{100*row["critical_satisfaction"]:.4f}%', f'{row["soc_delta"]:.4f}', f'{100*row["h2_delta_fraction"]:.4f}%', f'{row["h2_lower_bound"]["duration_h"]:.2f}']
            for year, row in yearly.items()
        ],
        [900, 1200, 1920, 1680, 1680, 1980],
    )
    add_paragraph(
        doc,
        "所有640个扩展候选均通过周期收敛、逐步功率守恒、H₂质量守恒、边界与最小稳定功率检查；全部因逐年LPSP、三年总LPSP、逐年关键负荷满足率及三年关键负荷满足率失败。",
    )

    new_page(doc)
    heading(doc, "四、年度能量预算与H₂下限诊断", 1)
    add_paragraph(doc, "下表给出扩展域LPSP最低配置的完整年度能量预算；数值仅用于解释容量准入失败，不用于寿命或经济外推。")
    add_table(
        doc,
        ["年", "负荷(MWh)", "原始PV(MWh)", "缩放PV(MWh)", "弃光(MWh)", "未供电(MWh)"],
        [
            [year, f'{row["load_energy_kwh"]/1000:.3f}', f'{row["pv_raw_energy_kwh"]/1000:.3f}', f'{row["pv_scaled_energy_kwh"]/1000:.3f}', f'{row["pv_curtailed_energy_kwh"]/1000:.3f}', f'{row["unserved_energy_kwh"]/1000:.3f}']
            for year, row in yearly.items()
        ],
        [720, 1680, 1680, 1800, 1680, 1800],
    )
    add_table(
        doc,
        ["年", "BAT充/放(MWh)", "EL输入(MWh)", "产氢(kg)", "FC输出(MWh)", "辅机(MWh)"],
        [
            [year, f'{row["battery_charge_energy_kwh"]/1000:.3f}/{row["battery_discharge_energy_kwh"]/1000:.3f}', f'{row["electrolyzer_input_energy_kwh"]/1000:.3f}', f'{row["h2_production_kg"]:.2f}', f'{row["fuel_cell_output_energy_kwh"]/1000:.3f}', f'{row["auxiliary_energy_kwh"]/1000:.3f}']
            for year, row in yearly.items()
        ],
        [720, 2040, 1560, 1440, 1800, 1800],
    )
    add_table(
        doc,
        ["年", "首次到达H₂下限", "持续时长(h)", "下限期间未供电(MWh)"],
        [
            [year, row["h2_lower_bound"]["first_timestamp"], f'{row["h2_lower_bound"]["duration_h"]:.2f}', f'{row["h2_lower_bound"]["unserved_energy_kwh"]/1000:.3f}']
            for year, row in yearly.items()
        ],
        [900, 3000, 2100, 3360],
    )
    add_paragraph(
        doc,
        f"最大逐步功率残差/额定功率={best['metrics']['max_power_residual_fraction']:.3e}，最大H₂质量残差/额定质量={best['metrics']['h2_residual_fraction']:.3e}；两者均低于1×10⁻⁶。",
    )

    new_page(doc)
    heading(doc, "五、准入规则与双层不确定性预注册", 1)
    add_table(
        doc,
        ["准入检查", "阈值", "扩展域结果"],
        [
            ["逐年及三年总LPSP", "≤1%", "全部失败"],
            ["逐年及三年关键负荷满足率", "≥99%", "全部失败"],
            ["三年块SOC首尾差", "≤5%", "全部通过"],
            ["三年块H₂首尾差/额定质量", "≤5%", "全部通过"],
            ["逐步功率与H₂质量残差", "≤1×10⁻⁶", "全部通过"],
            ["SOC/H₂边界与最小稳定功率", "逐步满足", "全部通过"],
        ],
        [3000, 2520, 3840],
    )
    add_paragraph(doc, "三年块首尾闭合是周期稳态门控；单年SOC/H₂偏差不设≤5%硬门槛，避免错误排除跨年季节储能。")
    heading(doc, "5.1 双层不确定性样本量与学习曲线停止条件", 2)
    add_table(
        doc,
        ["层级", "抽样设计", "检查点", "最低样本", "停止条件"],
        [
            ["外层参数", "嵌套Sobol；健康–性能/退化参数", "64/128/256/512", "256", "连续2检查点端点变化≤前一CI宽度5%"],
            ["内层时序", "季节分层4周移动块Bootstrap", "1000/2000/4000/8000", "4000", "连续2检查点激活概率变化≤0.02"],
        ],
        [1320, 2400, 1800, 1320, 2520],
        font_size=8.0,
    )
    add_paragraph(
        doc,
        "只有同时达到外层≥256、内层≥4000且两类条件连续满足2个检查点才停止；外层512、内层8000仍未满足则标记“蒙特卡洛未收敛”。当前容量准入为0，该设计仅完成配置与预注册，未产生任何路径不确定性结论。",
    )

    new_page(doc)
    heading(doc, "六、下游冻结与历史结果隔离", 1)
    add_paragraph(
        doc,
        "“未通过容量准入则停止下游物理论证”已写入配置与总门控结果。当前downstream_physical_inference_allowed=false，冻结对象为退化、六路径、Shapley、L4、经济和双层不确定性；现有代码测试只证明模块可执行或算术闭合，不证明可行系统中的物理有效性。",
    )
    add_table(
        doc,
        ["模块", "保留证据", "冻结内容"],
        [
            ["退化/EOL", "逐步封顶、事件接口与回归测试", "寿命量级、替换次数"],
            ["六路径", "统计门控代码与历史诊断输出", "路径激活物理结论"],
            ["Shapley", "路径对齐后的算术闭合", "归因机制有效性"],
            ["L4", "旧配置上的否决性测试", "控制优越性"],
            ["经济", "事件贴现算法", "NPC、LCOE与工程成本结论"],
            ["双层不确定性", "样本量与停止条件", "区间、概率与排序"],
        ],
        [1800, 3720, 3840],
    )
    add_paragraph(
        doc,
        f"历史目录annual、crossyear、economic、l4_comparison、lifetime、path_identification与shap中的{history['file_count']}个文件均已记录SHA-256并标记historical_invalid_pv_scaling、paper_citable=false；文件未移动、未删除，与07_results/rectification_v2严格隔离。",
    )
    heading(doc, "6.1 C5 2017年LPSP矛盾的统一处理", 2)
    add_table(
        doc,
        ["来源", "数值", "统一解释"],
        [
            ["F8跨年CSV", "31.2608013347%", "历史连续年诊断行；采用唯一CSV原值展示"],
            ["F4/F5经济摘要", "31.36%", "另一历史运行的四舍五入值；与F8不视为同一实验"],
            ["本轮旧C5-500三年块", "18.0958121077%", "按233 kW缩放与周期稳态重算；仍不准入"],
        ],
        [2160, 1800, 5400],
    )

    new_page(doc)
    heading(doc, "七、测试、断点续跑与可复现性", 1)
    add_paragraph(
        doc,
        "单元测试共34项全部通过。容量编排器按候选写入原子JSON检查点；输入哈希、代码哈希或配置哈希任一变化均使检查点失效。并行上限=min(请求值,8,逻辑核数−1)，候选种子为SHA-256(candidate_id)前8位十六进制整数。",
    )
    add_table(
        doc,
        ["编号", "名称", "输入", "期望/容差", "实际值", "状态"],
        [
            [item["check_id"], item["name"], item["input"], f'{item["expected"]}；tol={item["tolerance"]}', item["actual"], "PASS" if item["passed"] else "FAIL"]
            for item in s12["checks"]
        ],
        [900, 1560, 2100, 2280, 1560, 960],
        font_size=7.2,
    )
    add_paragraph(doc, "S12结构化日志：07_results/rectification_v2/tests/s12_regression.json；9项均明确记录名称、输入、期望值、实际值与容差。")
    add_table(
        doc,
        ["可复现字段", "值/位置"],
        [
            ["输入哈希", audit["input_hash"]],
            ["配置哈希", gate["domain_summaries"]["main"]["config_hash"]],
            ["容量代码哈希", gate["domain_summaries"]["main"]["code_hash"]],
            ["运行环境", f'Python {best["python_version"]}; NumPy {best["numpy_version"]}'],
            ["总门控", "07_results/rectification_v2/capacity_gate_summary.json"],
            ["评审台账", "09_docs/review_issue_register_2026-08-02.csv"],
        ],
        [2160, 7200],
        font_size=8.0,
    )

    new_page(doc)
    heading(doc, "八、76项评审的顺序处置状态", 1)
    add_paragraph(doc, "评审项按原编号R01–R76保留；完整逐项原文、阶段、证据路径与状态见评审台账。以下矩阵不改变依赖顺序。")
    add_table(
        doc,
        ["评审编号", "主题", "本轮状态", "理由/证据"],
        [
            ["R01–R07", "文档矛盾与表述", "已整改", "统一C5来源、八项门控、S12九项、候选路径措辞"],
            ["R08–R18", "PV/容量/年度能量", "已整改", "233 kW契约、完整年、三年块稳态、能量预算与联合搜索"],
            ["R19", "路径窗口继承年度状态", "冻结", "容量准入为0，未启动路径重辨识"],
            ["R20–R32", "FC/BAT/EL退化标定", "冻结", "等待3–5个真实准入配置"],
            ["R33–R37", "EOL事件", "代码证据保留；物理量冻结", "逐步封顶测试通过；失配系统结果不外推"],
            ["R38–R49", "六路径与双层传播", "冻结", "样本量/停止条件已预注册；未执行"],
            ["R50–R57", "Shapley归因", "冻结", "仅保留算术闭合代码证据"],
            ["R58–R64", "经济与L4", "冻结", "历史数值移入无效结果语境"],
            ["R65–R67", "测试与哈希", "已整改", "34项单测、S12结构化日志、三类哈希"],
            ["R68–R69", "Bootstrap种子与字段定义", "冻结", "容量门控失败，尚无新统计输出"],
            ["R70–R71", "合成数据与历史隔离", "已整改", "数据限制声明；42个历史文件状态清单"],
            ["R72–R76", "结论措辞", "已整改", "总判定置于页眉并冻结所有下游物理结论"],
        ],
        [1320, 2400, 1800, 3840],
        font_size=7.8,
    )

    heading(doc, "九、最终判定", 1)
    final = add_paragraph(
        doc,
        "主研究域90–150 kWp在必要能量条件下不可行；扩展域160–320 kWp虽有640组进入三年周期精筛，仍无配置达到LPSP≤1%和关键负荷满足率≥99%。因此尚不能冻结3–5个设计配置，阶段3–7不得启动物理论证；当前论文可报告的结论仅为“原设计域不可行，现有扩展域仍未建立连续年度可行基线”。",
        keep=True,
    )
    for run in final.runs:
        run.bold = True
    add_paragraph(
        doc,
        "代码级可保留项为PV输入契约、完整年审计、年际状态连续、三年块周期稳态、逐步守恒、必要能量筛选、并行断点续跑及准入门控。历史路径、寿命、L4和经济结果均不得回流为论文物理结论。",
    )

    doc.core_properties.title = "工作区全面总结报告｜阶段0–2正式实施版"
    doc.core_properties.subject = "PV输入契约、三年周期稳态容量门控与下游冻结"
    doc.core_properties.author = "Codex workspace rectification"
    doc.core_properties.keywords = "PV; H2; BESS; capacity admission; continuous three-year cycle"

    temporary = TARGET.with_suffix(".stage02.tmp.docx")
    doc.save(temporary)
    Document(temporary)
    temporary.replace(TARGET)
    print(TARGET)


if __name__ == "__main__":
    build()
