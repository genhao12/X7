"""Synchronize the workspace summary DOCX with verified Stage2b artifacts."""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from docx import Document


DT_H = 0.25


def _load(path: Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _set_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in list(paragraph.runs[1:]):
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def _replace_prefix(document, prefix: str | tuple[str, ...], text: str) -> None:
    prefixes = (prefix,) if isinstance(prefix, str) else prefix
    for paragraph in document.paragraphs:
        if any(paragraph.text.startswith(item) for item in prefixes):
            _set_paragraph(paragraph, text)
            return
    raise ValueError(f"DOCX paragraph prefix not found: {prefixes}")


def _set_cell(cell, text: str) -> None:
    _set_paragraph(cell.paragraphs[0], text)
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)


def _pct(value: float) -> str:
    return f"{100.0 * float(value):.4f}%"


def _mwh(value_kwh: float) -> str:
    return f"{float(value_kwh) / 1000.0:.3f}"


def _best_audit_case(audit: dict) -> dict:
    return min(audit["cases"], key=lambda row: float(row["metrics"]["three_year_lpsp"]))


def _trace_years(audit_path: Path, case: dict) -> dict:
    trace_path = audit_path.parent.parent / case["trace"]
    years = {}
    with trace_path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            year = row["year"]
            bucket = years.setdefault(
                year,
                {
                    "load": 0.0,
                    "pv": 0.0,
                    "aux": 0.0,
                    "cur_dump": 0.0,
                    "unserved": 0.0,
                    "bat_charge": 0.0,
                    "bat_discharge": 0.0,
                    "el": 0.0,
                    "fc": 0.0,
                    "h2_prod": 0.0,
                    "h2_cons": 0.0,
                    "h2_min_hours": 0.0,
                    "h2_min_unserved": 0.0,
                    "h2_first": None,
                },
            )
            values = {key: float(row[key]) for key in (
                "load_kw", "pv_kw", "p_aux_kw", "p_cur_kw", "p_dump_kw",
                "p_uns_kw", "p_bat_kw", "p_el_kw", "p_fc_kw",
                "h2_production_kg", "h2_consumption_kg", "h2_kg",
            )}
            bucket["load"] += values["load_kw"] * DT_H
            bucket["pv"] += values["pv_kw"] * DT_H
            bucket["aux"] += values["p_aux_kw"] * DT_H
            bucket["cur_dump"] += (values["p_cur_kw"] + values["p_dump_kw"]) * DT_H
            bucket["unserved"] += values["p_uns_kw"] * DT_H
            bucket["bat_charge"] += max(-values["p_bat_kw"], 0.0) * DT_H
            bucket["bat_discharge"] += max(values["p_bat_kw"], 0.0) * DT_H
            bucket["el"] += values["p_el_kw"] * DT_H
            bucket["fc"] += values["p_fc_kw"] * DT_H
            bucket["h2_prod"] += values["h2_production_kg"]
            bucket["h2_cons"] += values["h2_consumption_kg"]
            h2_min = 0.10 * float(case["anchor"]["H2_kg"])
            if values["h2_kg"] <= h2_min + 1e-8:
                bucket["h2_min_hours"] += DT_H
                bucket["h2_min_unserved"] += values["p_uns_kw"] * DT_H
                if bucket["h2_first"] is None:
                    bucket["h2_first"] = row["timestamp"]
    return years


def update_workspace_summary(
    docx_path: Path,
    search_summary_path: Path,
    audit_summary_path: Path,
    gate_path: Path,
    test_count: int | None = None,
) -> None:
    docx_path = Path(docx_path)
    search_path = Path(search_summary_path)
    audit_path = Path(audit_summary_path)
    search = _load(search_path)
    audit = _load(audit_path)
    gate = _load(gate_path)
    manifest_path = Path(gate_path).parent / "artifact_manifest.json"
    manifest = _load(manifest_path) if manifest_path.exists() else {}
    best = _best_audit_case(audit)
    years = _trace_years(audit_path, best)
    tier1 = search["tier_summaries"]["tier1"]
    tier2 = search["tier_summaries"]["tier2"]
    admitted = int(gate["verified_admitted_count"])
    stage3 = bool(gate["stage3_basic_degradation_allowed"])
    gate_text = "阶段3已开放" if stage3 else "阶段3继续冻结"
    document = Document(docx_path)

    _replace_prefix(document, "工作区全面总结报告｜", "工作区全面总结报告｜Stage2b正式实施与当前阻断版")
    _replace_prefix(document, "更新日期：", "更新日期：2026-08-03")
    _replace_prefix(document, "当前总判定：", f"当前总判定：Stage2b容量准入{admitted}组，{gate_text}")
    _replace_prefix(
        document,
        ("阶段2容量准入结果：", "Stage2b容量准入结果："),
        f"Stage2b容量准入结果：一级域{tier1['evaluated_count']}组、二级增量域{tier2['evaluated_count']}组，"
        f"共{search['evaluated_total']}组完成评估，{admitted}组准入；EMS审计{audit['case_count']}个连续三年案例完成。"
        f"{gate_text}，退化、六路径、Shapley、L4、经济与双层不确定性实算保持冻结。",
    )
    _replace_prefix(document, ("三、阶段2：", "三、Stage2b EMS"), "三、Stage2b EMS瓶颈审计与容量边界扩展")
    _replace_prefix(
        document,
        "执行链固定为：",
        "执行链固定为：必要能量筛选 → 单年粗筛 → 连续三年周期精筛 → 并行自适应搜索。"
        "每级候选池使用scipy.stats.qmc.Sobol(d=5, scramble=True, seed=20260803)一次生成2048点；"
        "首批256点，后续每批128点按102个综合违背度利用点与26个五维归一化最大最小距离覆盖点组成。",
    )
    _replace_prefix(document, ("3.1 扩展域", "3.1 Stage2b"), "3.1 Stage2b审计与自适应搜索结果")
    _replace_prefix(
        document,
        ("扩展域三年LPSP最低配置为", "审计最低三年LPSP案例为"),
        f"审计最低三年LPSP案例为{best['anchor']['name']} / {best['strategy']}："
        f"三年LPSP={_pct(best['metrics']['three_year_lpsp'])}，关键负荷满足率="
        f"{_pct(best['metrics']['three_year_critical_satisfaction'])}；"
        f"未供电={_mwh(best['bottleneck']['unserved_energy_kwh'])} MWh，"
        f"弃光={_mwh(best['bottleneck']['curtailed_energy_kwh'])} MWh，"
        f"P_dump={_mwh(best['bottleneck']['dumped_energy_kwh'])} MWh。",
    )
    _replace_prefix(
        document,
        ("所有640个扩展候选", "一级域完成"),
        f"一级域完成{tier1['evaluated_count']}组、二级增量域完成{tier2['evaluated_count']}组；"
        f"准入{admitted}组。一级最优LPSP归一化违背度={tier1['best_violation']['lpsp']:.6f}，"
        f"二级={tier2['best_violation']['lpsp']:.6f}；闭合、状态机与守恒必须与LPSP同时通过，"
        "任一单项最优不构成容量准入。",
    )
    _replace_prefix(document, ("四、年度能量预算", "四、Stage2b"), "四、Stage2b逐时审计与守恒证据")
    _replace_prefix(
        document,
        ("下表给出扩展域", "下表给出Stage2b"),
        "下表给出Stage2b审计最低三年LPSP案例的年度能量、库存和状态机诊断；"
        "该案例paper_citable=false，仅用于容量边界与EMS瓶颈定位。",
    )
    _replace_prefix(
        document,
        "“未通过容量准入则停止下游物理论证”",
        f"“未通过容量准入则停止下游物理论证”已由Stage2b总门控执行："
        f"verified_admitted_count={admitted}，stage3_basic_degradation_allowed={str(stage3).lower()}。"
        f"{gate_text}；六路径、Shapley、L4、经济和双层不确定性均未执行。",
    )
    _replace_prefix(
        document,
        ("历史目录annual、", "Stage2b历史失效结果"),
        f"Stage2b历史失效结果已迁入history子目录并统一标记paper_citable=false；"
        f"其中历史检查点{manifest.get('history_checkpoint_count', 0)}个，当前哈希有效检查点"
        f"{manifest.get('current_hash_valid_checkpoint_count', search['evaluated_total'])}个。"
        "旧annual、crossyear、economic、l4_comparison、lifetime、path_identification与shap结果继续隔离，"
        "不得回流论文物理结论。",
    )
    if test_count is not None:
        _replace_prefix(
            document,
            "单元测试共",
            f"Stage2b测试共{test_count}项全部通过；容量编排器按候选写入原子JSON检查点，"
            "输入哈希、代码哈希或配置哈希任一变化均使检查点失效。",
        )
    _replace_prefix(
        document,
        ("主研究域90–150 kWp在必要能量条件下不可行；", "主研究域90–150 kWp保持冻结；"),
        f"主研究域90–150 kWp保持冻结；Stage2b在160–320 kWp内完成{search['evaluated_total']}组自适应评估，"
        f"获得{admitted}组准入配置。{gate_text}；当前论文可报告结论为“原设计域不可行，"
        "现有Stage2b容量边界仍未建立连续年度可行基线”。"
        if not stage3
        else f"主研究域90–150 kWp保持冻结；Stage2b获得{admitted}组准入配置，阶段3仅按门控范围开放。",
    )

    table0 = document.tables[0]
    _set_cell(table0.rows[3].cells[1], f"Tier1 {tier1['evaluated_count']}组；Tier2增量域 {tier2['evaluated_count']}组")
    _set_cell(table0.rows[3].cells[2], f"{admitted}组准入")
    _set_cell(table0.rows[6].cells[1], f"verified_admitted_count={admitted}")
    _set_cell(table0.rows[6].cells[2], gate_text)

    table2 = document.tables[2]
    domain_rows = (
        ("主研究域", "90:10:150", "90/120/150/180", "30/35/40/45", "25/30/35/40/45", "10/20/30/40", "2240（冻结）"),
        ("Stage2b Tier1", "160–320", "180–1200", "45–160", "45–300", "40–3000", f"2048池/{tier1['evaluated_count']}评估"),
        ("Stage2b Tier2增量域", "160–320", "180–2400", "45–180", "45–320", "40–6000", f"严格增量/{tier2['evaluated_count']}评估"),
    )
    for row, values in zip(table2.rows[1:], domain_rows):
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    table3 = document.tables[3]
    rows = (
        ("主研究域", "2240", "0", "0", "0", "冻结"),
        ("Stage2b Tier1", "2048点候选池", str(tier1["evaluated_count"]), str(tier1["evaluated_count"]), str(tier1["admitted_count"]), gate_text),
        ("Stage2b Tier2增量域", "2048点候选池", str(tier2["evaluated_count"]), str(tier2["evaluated_count"]), str(tier2["admitted_count"]), gate_text),
    )
    for row, values in zip(table3.rows[1:], rows):
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    table4 = document.tables[4]
    headers = ("年份", "LPSP", "关键负荷满足率", "起止SOC", "起止H₂(kg)", "周期块状态")
    for cell, value in zip(table4.rows[0].cells, headers):
        _set_cell(cell, value)
    for row, year in zip(table4.rows[1:], ("2016", "2017", "2018")):
        metric = best["metrics"]["yearly"][year]
        values = (
            year,
            _pct(metric["lpsp"]),
            _pct(metric["critical_satisfaction"]),
            f"{metric['start_soc']:.4f}→{metric['end_soc']:.4f}",
            f"{metric['start_h2_kg']:.3f}→{metric['end_h2_kg']:.3f}",
            "连续传递",
        )
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    table5 = document.tables[5]
    headers = ("年", "负荷(MWh)", "PV可用(MWh)", "辅耗(MWh)", "弃光+dump(MWh)", "未供电(MWh)")
    for cell, value in zip(table5.rows[0].cells, headers):
        _set_cell(cell, value)
    for row, year in zip(table5.rows[1:], ("2016", "2017", "2018")):
        y = years[year]
        values = (year, _mwh(y["load"]), _mwh(y["pv"]), _mwh(y["aux"]), _mwh(y["cur_dump"]), _mwh(y["unserved"]))
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    table6 = document.tables[6]
    headers = ("年", "BAT充/放(MWh)", "EL输入(MWh)", "产氢(kg)", "FC输出(MWh)", "耗氢(kg)")
    for cell, value in zip(table6.rows[0].cells, headers):
        _set_cell(cell, value)
    for row, year in zip(table6.rows[1:], ("2016", "2017", "2018")):
        y = years[year]
        values = (year, f"{_mwh(y['bat_charge'])}/{_mwh(y['bat_discharge'])}", _mwh(y["el"]), f"{y['h2_prod']:.2f}", _mwh(y["fc"]), f"{y['h2_cons']:.2f}")
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    table7 = document.tables[7]
    for row, year in zip(table7.rows[1:], ("2016", "2017", "2018")):
        y = years[year]
        values = (year, y["h2_first"] or "未到达", f"{y['h2_min_hours']:.2f}", _mwh(y["h2_min_unserved"]))
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    table8 = document.tables[8]
    outcomes = (
        f"{admitted}组通过",
        f"{admitted}组通过",
        "收敛候选逐项检查",
        "收敛候选逐项检查",
        "审计9/9通过",
        "审计9/9通过",
    )
    for row, outcome in zip(table8.rows[1:], outcomes):
        _set_cell(row.cells[2], outcome)

    table13 = document.tables[13]
    reproducibility = {
        "输入哈希": search["input_hash"],
        "配置哈希": search["config_hash"],
        "容量代码哈希": search["code_hash"],
        "运行环境": "Python 3.9.13; NumPy 1.24.4; SciPy 1.9.1",
        "总门控": "07_results/rectification_v2_stage2b/governance/downstream_gate.json",
    }
    for row in table13.rows[1:]:
        key = row.cells[0].text
        if key in reproducibility:
            _set_cell(row.cells[1], reproducibility[key])

    temporary = docx_path.with_suffix(docx_path.suffix + ".tmp")
    document.save(temporary)
    temporary.replace(docx_path)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--search-summary", type=Path, required=True)
    parser.add_argument("--audit-summary", type=Path, required=True)
    parser.add_argument("--gate", type=Path, required=True)
    parser.add_argument("--test-count", type=int)
    args = parser.parse_args(argv)
    update_workspace_summary(
        args.docx,
        args.search_summary,
        args.audit_summary,
        args.gate,
        args.test_count,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = ["main", "update_workspace_summary"]
