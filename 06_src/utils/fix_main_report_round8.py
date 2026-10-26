# -*- coding: utf-8 -*-
"""Round-8 main report sync after gate opening (gate_passed=True)."""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\工作区全面总结报告.docx")


def set_para(doc, needle, new_text):
    for para in doc.paragraphs:
        if needle in para.text:
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
            print(f"[PARA] {needle[:20]!r} replaced")
            return True
    print(f"[WARN] para not found: {needle!r}")
    return False


def set_cell(doc, table_head, row_needle, col, new_text):
    for tb in doc.tables:
        if not tb.rows or not tb.rows[0].cells[0].text.startswith(table_head):
            continue
        for row in tb.rows:
            if row_needle in row.cells[0].text:
                row.cells[col].text = new_text
                print(f"[CELL] {row_needle!r} C{col} replaced")
                return True
        return False
    return False


def main() -> int:
    doc = docx.Document(str(DOC))

    # P6 global verdict
    for para in doc.paragraphs:
        if "均非论文可引用，须待门控通过" in para.text:
            new_p6 = (
                "Stage2b在160–320 kWp域内完成1792组评估，0组准入；能量平衡诊断定位到2016年9月季节性储能缺口为根因。"
                "Stage2c将PV域扩展至350–500 kWp，首批256组候选中149组通过全部准入门控，准入率58%。"
                "Stage3退化分析已完成（门控已通过：downstream_gate.json .5 gate_passed=True）"
                "（开环寿命、闭环Shapley C/P、L4否定性结果、辅耗敏感性、NPC/LCOE），退化寿命和经济指标已由 EFC 模型重算——"
                "EOL步内事件模型已完成21项TDD测试（149/149全量通过）；BAT循环损伤已升级为 EFC/能量吞吐模型"
                "（BAT_CYCLE_MODEL='efc'，280Ah 校准）；六路径已统一为 effect_pct 单位"
                "（D∈{0.3,0.5,0.7} 三水平），统一口径下 α₁/β₁ 量级大于跨设备路径（α₁ 最大 +24.33% "
                "@ hydrogen_dominant D=0.7；跨设备 |effect_pct|≤~3%）；单层LHS不确定性传播（128点）已完成（10.3b节），"
                "嵌套双层传播（外层Sobol×内层Bootstrap，256×4000）已完成（闭合残差≤1.6e-16）；"
                "成本敏感性分析已完成（10.3节，5×16=80 组）；保留局限为成本参数缺压缩/干燥/辅机/税费边界，"
                "成本边界未补全前相关数值不作工程经济评估依据。"
            )
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_p6
            print("[PARA] P6 replaced")
            break

    # P70
    set_para(doc, "在成本参数边界未补全前，本章数值不得作为物理论文结论或工程经济评估依据",
             "门控已通过（gate_passed=True），本章方法层结论可用于论文；但在成本参数边界未补全前，"
             "相关数值仍不作工程经济评估依据。")

    # P136/P137
    set_para(doc, "消融实验方法层计算已落地（l4_allowed=False，D_bat=0.5，4变体对比）",
             "消融实验已完成（l4_allowed=True，D_bat=0.5，4变体对比）：分别运行 R1单独、R2单独、R1+R2组合，"
             "并统计FC启停次数、FC运行小时数、BAT吞吐量（MWh/yr），以量化每条规则对各配置的独立贡献和交互效应。结果见下节。")
    set_para(doc, "消融实验方法层结果（l4_allowed=False，D_bat=0.5，4变体对比）：",
             "消融实验结果（l4_allowed=True，D_bat=0.5，4变体对比）：")

    # P156 nested double-layer progress
    set_para(doc, "③ 嵌套不确定性传播（执行中）",
             "③ 嵌套不确定性传播（已完成）：外层 Sobol 退化参数抽样 × 内层季节移动块 Bootstrap"
             "（256×4000，全方差闭合残差≤1.6e-16；认知占比 BAT 99.9%/FC 64-78%/EL 86-95%）。"
             "④ 成本敏感性与论文图表（F1–F5、T1–T5）已基于更新模型生成定稿。")

    # Table8 gate wording
    set_cell(doc, "评审编号", "R19", 2,
             "六路径方法层计算已落地（six_path_execution_started=True，门控已通过）")
    set_cell(doc, "评审编号", "R38–R49", 2,
             "六路径方法层计算已落地（单位已统一 effect_pct，3 水平）；嵌套双层传播已完成（256×4000）")
    set_cell(doc, "评审编号", "R38–R49", 3,
             "六路径单位已统一（effect_pct，D∈{0.3,0.5,0.7} 三水平）；统一单位下 α₁/β₁ 量级大于跨设备路径"
             "（α₁ 最大 +24.33% @ hydrogen_dominant D=0.7）；跨设备 |effect_pct|≤~3%，不能断言可忽略；"
             "嵌套双层传播已完成（256×4000）")
    set_cell(doc, "评审编号", "R50–R57", 2,
             "已完成（C/P, 4联盟；shapley_allowed=True）")
    set_cell(doc, "评审编号", "R58–R64", 2,
             "情景计算完成（economics_allowed=l4_allowed=True，门控已通过）")
    set_cell(doc, "评审编号", "R58–R64", 3,
             "NPC/LCOE 情景计算完成（缺成本边界）；L4 在 4/5 配置上增强耦合（l4_allowed=True，否定性结果如实报告），"
             "电池主导 phi_C +5.3%（非-5.3%）；消融实验已完成（R1/R2/R1+R2）")

    # Table24 gate wording
    set_cell(doc, "分析项", "闭环Shapley", 1,
             "六路径已执行（α₁/β₁ 见10.4；α₂/β₂/γ₁/γ₂ 见10.4b；单位已统一 effect_pct；shapley_allowed=True）")
    set_cell(doc, "分析项", "L4控制器", 1,
             "否定性结果+消融完成（l4_allowed=True，如实报告）")
    set_cell(doc, "分析项", "NPC/LCOE", 1,
             "情景计算；成本敏感性已完成（10.3节，5配置×16情景=80组，economics_allowed=True，门控已通过）；PEMEL为4/5配置第一敏感项")

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
