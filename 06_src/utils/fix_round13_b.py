"""Round13 补充修复：LCOE<8 / Table10 BAT / Table7 / 62 缺失点。"""
import docx
from pathlib import Path

ROOT = Path(r"d:\matlab\C")

# 1) 主报告 P106 LCOE<8 + 62 缺失点
p = ROOT / "09_docs" / "工作区全面总结报告.docx"
d = docx.Document(str(p))
for para in d.paragraphs:
    if "所有配置在参数不确定性范围内LCOE均<8" in para.text:
        full = "".join(r.text for r in para.runs)
        new = full.replace(
            "所有配置在参数不确定性范围内LCOE均<8 CNY/kWh（90%CI上界）",
            "所有配置经验 90% 区间上界 LCOE<8 CNY/kWh（来自 uncertainty_propagation.json p95）",
        )
        for r in para.runs:
            r.text = ""
        if para.runs:
            para.runs[0].text = new
        else:
            para.add_run(new)
        print("P106 LCOE<8 表述修正")
        break

# 62 缺失点插补敏感性
inserted = False
for para in d.paragraphs:
    if "数据契约" in para.text and "合成" in para.text:
        new_p = docx.oxml.OxmlElement("w:p")
        para._element.addnext(new_p)
        np_ = docx.text.paragraph.Paragraph(new_p, para._parent)
        np_.add_run(
            "62 个负荷缺失点插补敏感性说明：02_data_processed 中 64_TradeStreet CSV 含 62 个负荷缺失点（<0.2%），"
            "采用前后值线性填充；当前实现未给出插补敏感性分析（即未对比插补前后的容量准入/EOL 寿命/LCOE 差异），"
            "属已知数据限制——论文定稿前须补充对比基准情景（不插补 vs 插补）的容量准入门控与寿命估计差异，"
            "或在论文中明确声明该插补对结论的影响未量化。"
        )
        inserted = True
        print("62 缺失点插补敏感性段已插入")
        break
if not inserted:
    for para in d.paragraphs:
        if "10.1 退化模型与开环年损伤" in para.text:
            new_p = docx.oxml.OxmlElement("w:p")
            para._element.addnext(new_p)
            np_ = docx.text.paragraph.Paragraph(new_p, para._parent)
            np_.add_run(
                "62 个负荷缺失点插补敏感性说明：02_data_processed 中 64_TradeStreet CSV 含 62 个负荷缺失点（<0.2%），"
                "采用前后值线性填充；当前实现未给出插补敏感性分析，属已知数据限制——"
                "论文定稿前须补充对比基准情景（不插补 vs 插补）的容量准入门控与寿命估计差异。"
            )
            print("62 缺失点插补敏感性段已插入（兜底）")
            break
d.save(str(p))
print("main saved")

# 2) 汇总 Table10 R4 BAT 认知方差日历 + Table7 主要结果
p = ROOT / "09_docs" / "工作区运行结果汇总.docx"
d = docx.Document(str(p))
t10 = d.tables[10]
for row in t10.rows:
    for ci, c in enumerate(row.cells):
        if "BAT 认知占比 99.9%（日历项主导）" in c.text:
            c.text = c.text.replace(
                "BAT 认知占比 99.9%（日历项主导）",
                "BAT 认知占比 99.9%（BAT_K_CAL_H 固定 5%/年不进入 LHS，BAT 认知方差仅反映 EFC 循环参数 BAT_EFC_A_D/K 不确定性）",
            )
            print(f"Table10 C{ci} BAT 认知方差日历修正")
            break

for tb in d.tables:
    if len(tb.rows) > 1 and "npc_lite" in tb.rows[1].cells[1].text:
        for row in tb.rows:
            if row.cells[0].text.strip() == "主要结果":
                if "历史审计附录" not in row.cells[1].text:
                    new_text = (
                        "【历史审计附录】"
                        + row.cells[1].text
                        + "（legacy MVP 单配置 bat=300kWh；Stage2c 5 配置 EFC 经济：NPC 31.5~39.3M、LCOE 5.55~6.90，见 6.3/Table11）"
                    )
                    row.cells[1].text = new_text
                    print("Table7 主要结果加历史审计附录")
                break
        break
d.save(str(p))
print("summary saved")
