# -*- coding: utf-8 -*-
"""Round-8 ledger sync after gate opening (downstream_gate.json .5, gate_passed=True)."""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\统一结果台账.docx")

R10 = [
    "round8 门控通过（2026-08-04 17:1x）：downstream_gate.json(.5) gate_passed=True；BAT-EFC 循环损伤升级（149 测试）、EOL/经济重跑、六路径单位统一（3 水平）、嵌套双层 256×4000 传播（闭合残差≤1.6e-16）、成本敏感性与 F1-F5/T1-T5 全部完成",
    "门控字段：degradation_execution_started/six_path_execution_started/shapley_allowed/l4_allowed/economics_allowed/double_layer_uncertainty_allowed 全部 True；legacy_results_paper_citable 保持 False（旧 MVP 结果不可引用）",
    "证据：open_stage3_gate_round7.py 11 项核验全 PASS；double_layer_propagation.json n_outer=256/n_inner=4000 全配置闭合；run_manifest.json(.4) tests=140 函数/149 用例、bat_cycle_model=efc",
]

TRAILER = (
    "round8 门控通过核验（2026-08-04 17:1x）：downstream_gate.json(.5) gate_passed=True。"
    "执行顺序 ①-⑤ 全部完成——① BAT-EFC 循环损伤升级（closed_loop_lite.py 默认 efc，149 测试 PASS）；"
    "② EOL 寿命与 NPC/LCOE 重跑（degradation_summary.json economics）；"
    "③ 六路径单位统一（effect_pct，D∈{0.3,0.5,0.7} 三水平）；"
    "④ 嵌套双层传播 256×4000（全方差闭合残差≤1.6e-16，认知占比 BAT 99.9%/FC 64-78%/EL 86-95%）；"
    "⑤ 成本敏感性（5×16=80 组）与 F1-F5/T1-T5 定稿。"
    "legacy_results_paper_citable 保持 False（旧 MVP 代理模型结果仍不可引用）。"
)


def main() -> int:
    doc = docx.Document(str(DOC))

    # ---- paragraphs ----
    for para in doc.paragraphs:
        t = para.text
        if t.startswith("4　门控通过后的执行顺序（当前未启动）"):
            for run in para.runs:
                run.text = ""
            para.runs[0].text = "4　执行顺序与门控通过记录（2026-08-04 已通过）"
            print("P7 updated")
        elif t.startswith("一致性门控目标："):
            for run in para.runs:
                run.text = ""
            para.runs[0].text = (
                "一致性门控结果（2026-08-04 17:1x）：BAT-EFC 循环损伤升级（149 测试）、EOL/经济重跑、"
                "六路径单位统一（3 水平）、嵌套双层 256×4000 传播（闭合残差≤1.6e-16）、"
                "成本敏感性与论文图表（F1–F5/T1–T5）均已执行完成；downstream_gate.json(.5) gate_passed=True。执行记录："
            )
            print("P8 updated")
        elif t.startswith("① BAT 循环损伤模型升级"):
            for run in para.runs:
                run.text = ""
            para.runs[0].text = "① BAT 循环损伤模型升级（已完成）：BATKCYC_H 运行小时代理 → EFC/能量吞吐模型（280Ah 校准，149 测试）"
            print("P9 updated")
        elif t.startswith("② 重跑 EOL 寿命"):
            for run in para.runs:
                run.text = ""
            para.runs[0].text = "② 重跑 EOL 寿命（EOL 步内事件模型）与 NPC/LCOE（M4 经济）（已完成，degradation_summary.json economics）"
            print("P10 updated")
        elif t.startswith("③ 统一六路径效应单位"):
            for run in para.runs:
                run.text = ""
            para.runs[0].text = "③ 统一六路径效应单位（effect_pct 归一化）及三退化水平（D∈{0.3,0.5,0.7}）（已完成）"
            print("P11 updated")
        elif t.startswith("④ 启动外层 Sobol"):
            for run in para.runs:
                run.text = ""
            para.runs[0].text = "④ 启动外层 Sobol × 内层季节移动块 Bootstrap 嵌套双层传播（已完成：256×4000，闭合残差≤1.6e-16）"
            print("P12 updated")
        elif t.startswith("⑤ 成本敏感性与论文图表"):
            for run in para.runs:
                run.text = ""
            para.runs[0].text = "⑤ 成本敏感性与论文图表（F1–F5、T1–T5）基于更新模型重算并定稿（已完成）"
            print("P13 updated")

    # ---- Table1 gate column ----
    t1 = doc.tables[1]
    for row in t1.rows:
        c0 = row.cells[0].text.strip()
        if c0.startswith("S2"):
            row.cells[6].text = "149 准入，5 冻结代表；阶段3 门控已通过（gate_passed=True）"
        elif c0.startswith("S8"):
            row.cells[6].text = "shapley_allowed=True（门控已通过）"
        elif c0.startswith("S9"):
            row.cells[3].text = "run_double_layer_propagation.py（真实嵌套：256 Sobol × 4000 Bootstrap，单年 2017 调度）"
            row.cells[6].text = "double_layer_uncertainty_allowed=True（门控已通过，256×4000）"
            row.cells[7].text = (
                "256 Sobol×4000 Bootstrap 真实嵌套，全方差闭合残差≤1.6e-16；认知占比 BAT 99.9%/FC 64-78%/EL 86-95%；"
                "inner_method=season_stratified_4week_moving_block_bootstrap_on_dispatch"
            )
        elif c0.startswith("S10"):
            row.cells[6].text = "economics_allowed=True（门控已通过）"
        elif c0.startswith("S11"):
            row.cells[6].text = "l4_allowed=True（门控已通过；结果仍为否定性，如实报告）"
    print("Table1 gate column updated")

    # ---- Table2 gate flags ----
    t2 = doc.tables[2]
    for row in t2.rows:
        key = row.cells[0].text.strip()
        if key == "stage3_basic_degradation_allowed":
            row.cells[1].text = "True（门控已通过：正式退化执行已完成，degradation_execution_started=True）"
        elif key == "robust_six_path_planning_allowed":
            row.cells[1].text = "True（six_path_execution_started=True）"
        elif key == "shapley_allowed":
            row.cells[1].text = "True"
        elif key == "l4_allowed":
            row.cells[1].text = "True（结果仍为否定性，如实报告）"
        elif key == "economics_allowed":
            row.cells[1].text = "True"
        elif key == "double_layer_uncertainty_allowed":
            row.cells[1].text = "True（256×4000 真实嵌套已执行）"
        elif key == "legacy_results_paper_citable":
            row.cells[1].text = "False（永久：旧 MVP 代理模型结果不可引用）"
        elif key == "gate_reason":
            row.cells[1].text = (
                "BAT-EFC 完成（149 测试）；EOL/经济重跑；六路径单位统一（3 水平）；嵌套双层 256×4000 闭合 PASS；"
                "F1-F5/T1-T5 生成；执行/论文标志 True；legacy_results_paper_citable=False"
            )
        elif key == "gate_passed":
            row.cells[1].text = "True（门控已通过：downstream_gate.json .5，2026-08-04 17:1x）"
    print("Table2 gate flags updated")

    # ---- Table3 append R10 ----
    t3 = doc.tables[3]
    nr = t3.add_row()
    for ci, txt in enumerate(R10):
        nr.cells[ci].text = txt
    print("Table3 appended row:", len(t3.rows) - 1)

    # ---- trailer ----
    doc.add_paragraph(TRAILER)
    print("trailer appended")

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
