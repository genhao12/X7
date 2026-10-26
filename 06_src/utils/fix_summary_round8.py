# -*- coding: utf-8 -*-
"""Round-8 summary report sync after gate opening (gate_passed=True)."""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\工作区运行结果汇总.docx")


def set_para(doc, needle, new_text):
    for para in doc.paragraphs:
        if needle in para.text:
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
            print(f"[PARA] {needle[:22]!r} replaced")
            return True
    print(f"[WARN] para not found: {needle!r}")
    return False


def main() -> int:
    doc = docx.Document(str(DOC))

    # P33 gate status paragraph
    set_para(doc, "【门控状态】per downstream_gate.json",
             "设计目的：在 5 个冻结代表配置上，完成单层参数不确定性传播、成本敏感性、跨设备六路径耦合分解，"
             "并为嵌套双层传播（外层 Sobol × 内层 Bootstrap）奠定基础。"
             "【门控状态】per downstream_gate.json(.5)：gate_passed=True（2026-08-04 门控已通过）；"
             "shapley_allowed=l4_allowed=economics_allowed=double_layer_uncertainty_allowed=True；"
             "legacy_results_paper_citable=False（旧 MVP 代理模型结果不可引用）。"
             "本节约述的 Stage3 分析（单层 LHS、成本敏感性、六路径、嵌套双层、L4 消融）均为已完成执行，"
             "其中 L4 为否定性结果（如实报告）。")

    # Table9 LHS status
    for tb in doc.tables:
        if len(tb.rows) < 2 or tb.rows[1].cells[0].text.strip() != "脚本":
            continue
        r1 = tb.rows[1].cells[1].text
        if "run_degradation_uncertainty" in r1:
            for row in tb.rows:
                if row.cells[0].text.strip() == "是否通过":
                    row.cells[1].text = "通过（128 点 LHS 经验 90% 不确定性区间已给出；闭合性自检）"
                    print("Table9 PASS updated")
        elif "run_cost_sensitivity" in r1:
            for row in tb.rows:
                if row.cells[0].text.strip() == "是否通过":
                    row.cells[1].text = "通过（基线复现 Stage3 相对偏差 <1e-9；单价单调；全设备包络线性；economics_allowed=True）"
                    print("Table11 PASS updated")
        elif "run_cross_device_paths" in r1:
            for row in tb.rows:
                if row.cells[0].text.strip() == "是否通过":
                    row.cells[1].text = "通过（六路径覆盖 6/6、七健康 7/7、数值泄漏残差 0；单位已统一 effect_pct；shapley_allowed=True）"
                    print("Table12 PASS updated")

    # Table10 (double-layer) status: 32x25 -> 256x4000
    for tb in doc.tables:
        if len(tb.rows) < 2 or tb.rows[1].cells[0].text.strip() != "脚本":
            continue
        r1 = tb.rows[1].cells[1].text
        if "run_double_layer_propagation" in r1:
            for row in tb.rows:
                label = row.cells[0].text.strip()
                if label == "输入数据":
                    row.cells[1].text = (
                        "5 配置；外层 256 Sobol × 内层 4000 季节移动块 Bootstrap（真实嵌套规格，"
                        "minimum_inner_samples=4000 已满足）；单年 2017 调度"
                    )
                elif label == "主要结果":
                    row.cells[1].text = (
                        "BAT 认知占比 99.9%（日历项主导）；FC 认知 64-78%（偶然显著）；EL 认知 86-95%；"
                        "全方差闭合残差 ≤1.6e-16"
                    )
                elif label == "当前状态":
                    row.cells[1].text = "嵌套双层传播已完成（256×4000）；double_layer_uncertainty_allowed=True"
                elif label == "是否通过":
                    row.cells[1].text = "通过（全方差闭合残差 ≤1e-9；256×4000 真实嵌套）"
            print("Table10 double-layer updated")

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
