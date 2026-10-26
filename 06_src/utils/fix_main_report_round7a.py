# -*- coding: utf-8 -*-
"""Round-7a pre-update of 工作区全面总结报告.docx (parallel with ④):
1) test count 127/136 -> 140/149 (EFC 13)
2) units unified statements across 10.4/10.4b conclusions
3) EFC model info (BAT_CYCLE_MODEL='efc' default) & progress updates
"""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\工作区全面总结报告.docx")


def set_para(doc, needle, new_text):
    for para in doc.paragraphs:
        if needle in para.text:
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
            print(f"[PARA] {needle[:22]!r} replaced")
            return True
    print(f"[WARN] para not found: {needle!r}")
    return False


def set_cell(doc, table_head, row_needle, col, new_text):
    for tb in doc.tables:
        if not tb.rows:
            continue
        if not tb.rows[0].cells[0].text.startswith(table_head):
            continue
        for row in tb.rows:
            if row_needle in row.cells[0].text:
                row.cells[col].text = new_text
                print(f"[CELL] {table_head!r}/{row_needle!r} C{col} replaced")
                return True
        return False
    return False


def main() -> int:
    doc = docx.Document(str(DOC))

    # ---- 1) P61 test count ----
    set_para(doc, "测试构成（127 个测试函数经 @parametrize 展开为 136 个执行用例，全部通过）",
             "测试构成（140 个测试函数经 @parametrize 展开为 149 个执行用例，全部通过）："
             "Stage2b（8个文件）43 函数/50 用例 + unit/（12个文件）32 函数/34 用例 + "
             "Stage-3（1个文件）12 用例 = 96 用例基础；另有4个专项文件共 40 用例"
             "（Shapley C/P归因 6 + 退化时间单位 6 + EOL系统状态 21 + L4控制器 7）；"
             "再新增 EFC 循环损伤 13 用例（test_bat_efc_cycle_model.py）；"
             "合计 96 + 40 + 13 = 149 用例。容量编排器按候选写入原子JSON检查点，"
             "输入哈希、代码哈希或配置哈希任一变化均使检查点失效。")

    # ---- 2) P70 limitations ----
    set_para(doc, "⚠ 本章所有退化寿命和经济指标均为情景计算结果",
             "⚠ 本章所有退化寿命和经济指标均为情景计算结果，使用 MVP 退化模型（BAT 循环损伤已升级为 "
             "EFC/能量吞吐模型，BAT_CYCLE_MODEL='efc'，280Ah 校准 a_d=5.0714e-4 D/EFC）和文献基线成本参数。"
             "退化参数已完成单位转换（per-step→per-hour），EOL步内事件模型、更换复位和状态传递已完成21项"
             "TDD测试（149/149全量通过，见10.1节），但仍保留以下关键局限："
             "(1) 成本参数缺少价格年、压缩/干燥/辅机、税费边界；"
             "(2) 嵌套双层传播（外层Sobol×内层季节移动块Bootstrap，minimum_inner_samples=4000）执行中，"
             "参数不确定性区间仍宽（FC寿命CI达±factor 2.3）；"
             "(3) “90%区间”为128点LHS经验区间（未声明参数概率分布，非统计置信区间）。"
             "在成本参数边界未补全前，本章数值不得作为物理论文结论或工程经济评估依据。")

    # ---- 3) P78 ----
    set_para(doc, "这些数值依赖于BATKCYC_H运行小时代理",
             "这些数值基于 EFC 循环损伤模型（280Ah 校准）与 FC/EL 启停参数的不确定性区间（±50-100%），"
             "在参数不确定性传播完成前仅为情景参考。")

    # ---- 4) P106 ----
    set_para(doc, "范围限制：本节（10.4）仅覆盖",
             "范围限制：本节（10.4）仅覆盖BAT→FC（α₁）和BAT→EL（β₁）两条路径。"
             "跨设备路径α₂(EL→FC)、β₂(FC→EL)、γ₁(FC→BAT)、γ₂(EL→BAT)已在10.4b节实现并报告。"
             "cross_device_paths.json 已为全部六路径统一提供 effect_pct（相对目标设备自身退化的百分比贡献，"
             "D∈{0.3,0.5,0.7} 三水平），可在统一口径下比较量级；10.4 节表格另附 α₁/β₁ 的 ΔD/yr 绝对单位数值作参考。"
             "\"零闭合误差\"是2归因Shapley恒等闭合的数学性质（phiC + phiP ≡ v(11) - v(00)），不证明物理归因有效性。")

    # ---- 5) P107 ----
    set_para(doc, "效应单位：表中数值为ΔD/yr",
             "效应单位：10.4 节表中数值为 ΔD/yr（无量纲损伤增量/年），基于2017年单年调度，非三年平均，"
             "不确定性区间未计算；六路径统一比较请使用 cross_device_paths.json 的 effect_pct（百分比贡献，"
             "D∈{0.3,0.5,0.7} 三水平，单位已统一）。")

    # ---- 6) P121 ----
    set_para(doc, "与10.4a节α₁/β₁的ΔD/yr绝对单位不同",
             "对每个代表配置在D_source∈{0.3, 0.5, 0.7}下运行反事实调度：源设备使用退化参数"
             "（效率×(1-0.10D)、功率×(1-0.05D)），测量目标设备损伤变化。"
             "cross_device_paths.json 已为全部六路径统一提供 effect_pct（相对目标设备自身退化的百分比贡献，"
             "D∈{0.3,0.5,0.7} 三水平），六路径可在统一口径下比较量级；10.4 节 α₁/β₁ 另附 ΔD/yr 绝对单位数值。")

    # ---- 7) P125 ----
    set_para(doc, "D_source=0.5 下 γ₁/γ₂ 百分比贡献 <0.1%",
             "γ₁/γ₂（FC/EL→BAT）百分比贡献在 D∈{0.3,0.5,0.7} 三水平下均 <0.1%（effect_pct，相对 BAT 自身退化），"
             "FC/EL 退化对 BAT 损伤影响很小；全水平数据已补齐（cross_device_paths.json），"
             "但跨设备路径 |effect_pct| 最大仍 ≤~3%（α₂ −1.71%~+2.97%），不能据此断言跨设备路径“可忽略”。")

    # ---- 8) P126 ----
    set_para(doc, "α₂（EL→FC）和β₂（FC→EL）效应小",
             "α₂（EL→FC）和β₂（FC→EL）效应小（effect_pct 多数 |·|<2%）且方向配置依赖"
             "（cross_device_paths.json 已统一 effect_pct 口径，D∈{0.3,0.5,0.7} 三水平）：")

    # ---- 9) P129 ----
    set_para(doc, "与α₁/β₁（BAT→FC/EL，ΔD/yr绝对单位）的关系",
             "与α₁/β₁（BAT→FC/EL）的关系：cross_device_paths.json 已统一六路径单位为 effect_pct"
             "（相对目标设备自身退化的百分比贡献，D∈{0.3,0.5,0.7} 三水平）。统一口径下 α₁/β₁ 量级大于跨设备路径："
             "α₁ 最大 +24.33% @ hydrogen_dominant D=0.7、β₁ 最大 +10.47% @ critical D=0.7；"
             "跨设备路径 |effect_pct| ≤ ~3%（α₂ −1.71%~+2.97%，β₂/γ₁/γ₂ <1.5%）。")

    # ---- 10) P131 conclusion ----
    set_para(doc, "结论：六条退化耦合路径均非零",
             "结论：六条退化耦合路径均非零。cross_device_paths.json 已统一六路径单位为 effect_pct"
             "（相对目标设备自身退化的百分比贡献，D∈{0.3,0.5,0.7} 三水平）；统一口径下 "
             "α₁(BAT→FC)和β₁(BAT→EL) 的量级大于四条跨设备路径（α₁ 最大 +24.33% @ hydrogen_dominant D=0.7；"
             "跨设备 |effect_pct| ≤ ~3%）。跨设备路径方向配置依赖、量级较小，但不能断言“可忽略”或“仅 2 条激活”。")

    # ---- 11) P151/P154/P155/P156 progress ----
    set_para(doc, "下一步工作顺序（已进入Stage3模型闭合；嵌套双层不确定性传播待启动）",
             "执行顺序进度（BAT-EFC 循环损伤升级已完成；EOL/经济已重跑；六路径单位已统一；嵌套双层传播执行中）：")
    set_para(doc, "① BAT循环损伤模型升级：将BATKCYC_H运行小时代理改为EFC/能量吞吐",
             "① BAT循环损伤模型升级（已完成）：BATKCYC_H 运行小时代理已改为 EFC/能量吞吐模型"
             "（closed_loop_lite.py 默认 BAT_CYCLE_MODEL='efc'，rainflow+Woehler 深度加权，280Ah 校准 "
             "a_d=5.0714e-4 D/EFC），并已重跑 EOL 步内事件模型与经济评价（NPC/LCOE，见 degradation_summary.json economics）。")
    set_para(doc, "② 补齐六路径：在D_source∈{0.3,0.5,0.7}三水平下补全",
             "② 统一六路径单位（已完成）：cross_device_paths.json 已统一为 effect_pct（百分比贡献，"
             "D∈{0.3,0.5,0.7} 三水平）；统一口径下 α₁/β₁ 量级大于跨设备路径，"
             "“全水平<3%”的未证实断言已移除。")
    set_para(doc, "③ 嵌套不确定性传播：外层Sobol退化参数抽样",
             "③ 嵌套不确定性传播（执行中）：外层 Sobol 退化参数抽样 × 内层季节移动块 Bootstrap"
             "（minimum_inner_samples=4000）。④ 成本敏感性与论文图表（F1–F5、T1–T5）已生成，"
             "待嵌套双层完成后基于更新模型复核定稿。")

    # ---- 12) Table8 R38-R49 ----
    set_cell(doc, "评审编号", "R38–R49", 2,
             "六路径方法层计算已落地（单位已统一 effect_pct，3 水平）；嵌套双层传播执行中")
    set_cell(doc, "评审编号", "R38–R49", 3,
             "六路径单位已统一（effect_pct，D∈{0.3,0.5,0.7} 三水平）；统一单位下 α₁/β₁ 量级大于跨设备路径"
             "（α₁ 最大 +24.33% @ hydrogen_dominant D=0.7）；跨设备 |effect_pct|≤~3%，不能断言可忽略；"
             "嵌套双层传播执行中")

    # ---- 13) Table24 R4 ----
    set_cell(doc, "分析项", "闭环Shapley", 1,
             "六路径已执行（α₁/β₁ 见10.4；α₂/β₂/γ₁/γ₂ 见10.4b；单位已统一 effect_pct；shapley_allowed=False）")
    set_cell(doc, "分析项", "闭环Shapley", 2,
             "效应均非零；单位已统一（effect_pct，D∈{0.3,0.5,0.7}），α₁/β₁ 量级大于跨设备路径"
             "（α₁ 最大 +24.33%）；跨设备 |effect_pct|≤~3%，不能断言可忽略")

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
