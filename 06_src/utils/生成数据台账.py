# -*- coding: utf-8 -*-
"""生成数据来源与可信度台账Word文档"""
import sys, os
sys.path.insert(0, r"D:\matlab\C\运行脚本")
from 数据台账_data import SECTIONS
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
st = doc.styles['Normal']
st.font.name = '宋体'
st.font.size = Pt(10.5)
st._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

def set_bg(cell, c):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), c)
    pr.append(shd)

def H(text, lvl=1):
    h = doc.add_heading(text, level=lvl)
    for r in h.runs:
        r.font.name = '黑体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        r.font.size = Pt(18 if lvl == 0 else (13 if lvl == 1 else 11))
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

def P(text, bold=False, size=10.5, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = '宋体'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    r.font.size = Pt(size)
    r.bold = bold
    if align:
        p.alignment = align
    return p

def make_table(headers, rows, cw=None, hbg='1F3A5F'):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    hc = t.rows[0].cells
    for i, h in enumerate(headers):
        hc[i].text = ''
        p = hc[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.name = '宋体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        set_bg(hc[i], hbg)
        hc[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for rd in rows:
        c = t.add_row().cells
        for i, v in enumerate(rd):
            c[i].text = ''
            p = c[i].paragraphs[0]
            r = p.add_run(str(v))
            r.font.size = Pt(8)
            r.font.name = '宋体'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            c[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if cw:
        for i, w in enumerate(cw):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

HD = ['编号', '数据来源(DOI/名称/URL)', '数据内容', '用途', '可信度']
CW = [1.0, 5.4, 4.8, 3.5, 2.1]

# ===== 封面 =====
H('离网 PV-H2-LFP 微电网方法与情景研究\n数据来源与可信度台账', 0)
P('文档版本: 2026-07-29  |  论文定位: 方法与情景研究', bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()

H('0. 文档说明', 1)
P('本文档对当前研究已核查的全部数据进行归类整理,逐项标注:'
  '(1) 数据来源(DOI/名称/URL,可直接用于论文引用);'
  '(2) 数据内容(属于曲线/数据集/模型/额定点等);'
  '(3) 用途(用于论文哪个部分);'
  '(4) 可信度等级。'
  '研究定位为"基于公开数据、文献机理模型与显式不确定性情景的离网微电网容量优化与退化感知能量管理研究",'
  '所有设备性能、退化与经济性结论均表述为情景结果,而非真实工程设备或站点的确定性结论。')
P('')
P('可信度分级定义:', bold=True)
make_table(['等级', '定义'], [
    ['高', '公开原始数据集或经核验的厂家数据表,化学体系/功率口径/边界明确,可直接标定或作为强证据'],
    ['中', '文献机理模型、综述范围或厂家额定点,结构/方法可靠但数值不可直接迁移,需以情景区间使用'],
    ['低', '仅提供方法参考或跨规模/跨场景资料,数值不可用于本项目标定,仅作敏感性或背景论述'],
    ['不可用', '文件命名与正文不符、化学体系不匹配或口径不满足项目判据,不得进入证据链'],
], cw=[2.0, 15.0])

# ===== 各章节 =====
for sec in SECTIONS:
    doc.add_page_break()
    H(sec['title'], 1)
    if sec.get('intro'):
        P(sec['intro'], size=9.5)
    make_table(HD, sec['rows'], cw=CW)

# ===== 总结 =====
doc.add_page_break()
H('11. 数据缺口与论文可信度总评', 1)

P('11.1 当前数据可支撑的论文目标', bold=True, size=11)
P('(1) 15min离网EMS与容量协同优化:PV、LFP、PEMEL、储氢、PEMFC功率平衡;SOC、氢库存、0.05%/d泄漏、'
  '弃光、失供电;LPSP<=1%硬可靠性约束;词典序目标(可靠性->安全边界->NPC/LCOE->损伤/更换->弃光)。')
P('(2) 储能协同机制:电池短时波动平滑与日内移峰;PEMEL吸收富余PV;氢储能跨日保障;PEMFC低光照/低SOC时供电。')
P('(3) 文献驱动退化耦合情景:LFP日历(Naumann)+循环(Severson)分项标定;PEMEL OCV暴露/启停/负荷波动归一化损伤;'
  'PEMFC启停/低负荷/高负荷/负荷变化归一化损伤。')
P('(4) 20年技术经济区间与稳健性:统一币种/价格年/折现假设下CAPEX/O&M/替换成本区间;'
  '不同效率/成本/寿命/泄漏情景下NPC/LCOE范围与关键不确定性识别。')
P('(5) 方法学贡献:数据来源-参数边界-可迁移性-情景不确定性审计框架;'
  '逐点可溯源合成PV-建筑负荷层测试;真实同期原生数据缺失时结论变化报告。')

P('')
P('11.2 论文可信度判断', bold=True, size=11)
P('若严格按方法与情景研究定位写作,论文整体可信度中等偏高(约7/10)。', bold=True)
P('较强部分:系统建模、优化框架、可靠性约束、情景设计、LFP退化数据来源、可复现性和不确定性披露。')
P('较弱但可接受部分:PEMFC/PEMEL的设备级效率、动态和绝对退化率,需通过情景范围而非单一参数处理。')

P('')
P('11.3 使论文可信度成立的四个硬要求', bold=True, size=11)
reqs = [
    ('所有非实测设备参数都要标注证据等级', '如厂家额定点、文献机理、公开电芯数据、研究假设'),
    ('对关键缺口做区间而非单点假设', '特别是PEMFC氢耗、PEMEL部分负荷效率、BESS整包效率和成本'),
    ('报告敏感性与稳健性,不只报告一个最优解', '不同效率/退化/成本场景下,容量解是否稳定,LPSP是否仍满足'),
    ('避免把合成重建序列称为原生站点实测数据', '应称为具有逐点供体溯源的合成评价序列或重构的评价层'),
]
for i, (t, d) in enumerate(reqs, 1):
    P(f'  {i}. {t}: {d}', size=10)

P('')
P('11.4 仍需补齐的关键数据', bold=True, size=11)
priority = [
    '完整同期的原生PV-负荷-BESS实际年(决定研究能否升级为真实站点长期结论)',
    '目标PEMFC的净DC功率-氢耗多负荷点与动态数据(决定氢电耦合效率与可靠性结论可信度)',
    'Nel H4的分负荷、待机、启停与动态资料(决定PEMEL模型精度)',
    'BESS整包SOC/温度/功率约束和效率图(决定15min EMS是否有设备边界依据)',
    '30bar固定式储氢罐规格(决定储氢CAPEX精度)',
]
for i, item in enumerate(priority, 1):
    P(f'  {i}. {item}', size=10)

# 保存
out = r'D:\matlab\C\09_docs\06_项目梳理\09_数据来源与可信度台账.docx'
doc.save(out)
print(f'文档已生成: {out}')
print(f'共 {len(SECTIONS)} 个数据分类章节')
total = sum(len(s['rows']) for s in SECTIONS)
print(f'共 {total} 条数据记录')
