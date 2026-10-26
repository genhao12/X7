# -*- coding: utf-8 -*-
"""Round-4 final verification across ledger / summary / main report / governance JSON."""
from __future__ import annotations

import json
from pathlib import Path

import docx

ROOT = Path(r"d:\matlab\C")

docs = {
    "ledger": ROOT / "09_docs" / "统一结果台账.docx",
    "summary": ROOT / "09_docs" / "工作区运行结果汇总.docx",
    "main": ROOT / "09_docs" / "工作区全面总结报告.docx",
}

# collect all text (paragraphs + table cells) per doc
def collect(path: Path) -> list[str]:
    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                out.append(c.text)
    return out

texts = {k: "\n".join(v) for k, v in {k: collect(p) for k, p in docs.items()}.items()}

checks = [
    # (label, needle, docs that SHOULD contain it, docs that should NOT)
    ("2016 leap-year points (ledger/main)", "35136", ["ledger", "main"], []),
    ("2016 leap-year points (summary, thousands sep)", "35,136", ["summary"], []),
    ("cost groups 5x16=80", "80 组", ["ledger", "summary"], []),
    ("cost groups 80 in main", "5配置×16情景=80组", ["main"], []),
    ("stage2b 43 fn / 50 cases", "43 函数/50 用例", ["ledger", "main"], []),
    ("127 fn -> 136 cases (ledger/main)", "127", ["ledger", "main"], []),
    ("127 fn -> 136 cases (summary)", "127 个测试函数", ["summary"], []),
    ("six paths non-zero claim kept", "均非零", ["ledger", "summary"], []),
    ("negligible claim REMOVED (main)", "γ₁和γ₂效应可忽略", [], ["main"]),
    ("only-2-active heading REMOVED (main)", "仅α₁/β₁）", [], ["main"]),
    ("plan_only semantics (ledger)", "plan_only", ["ledger"], []),
    ("gate_passed False (ledger T2)", "False（门控未通过", ["ledger"], []),
    ("method-layer PASS (summary)", "方法层通过", ["summary"], []),
    ("governance hashes (main)", "11ed37be41737aaba1558b2eed366457d50724c73e9ac7094ea2f9afd249965f", ["main"], []),
]

all_ok = True
for label, needle, should, should_not in checks:
    for dk in should:
        ok = needle in texts[dk]
        print(f"[{'OK' if ok else 'FAIL'}] {label} -> {dk} contains {needle!r}: {ok}")
        all_ok &= ok
    for dk in should_not:
        bad = needle in texts[dk]
        print(f"[{'OK' if not bad else 'FAIL'}] {label} -> {dk} MUST NOT contain {needle!r}: {not bad}")
        all_ok &= (not bad)

# governance JSON checks
gov = ROOT / "07_results" / "rectification_v2_stage2c" / "governance"
gate = json.loads((gov / "downstream_gate.json").read_text(encoding="utf-8"))
manifest = json.loads((gov / "run_manifest.json").read_text(encoding="utf-8"))
hashrec = json.loads((gov / "hash_record.json").read_text(encoding="utf-8"))

json_checks = [
    ("gate.gate_passed False", gate["gate_passed"] is False),
    ("gate.stage3_basic_degradation_allowed True (plan-level)", gate["stage3_basic_degradation_allowed"] is True),
    ("gate.degradation_execution_started False", gate["degradation_execution_started"] is False),
    ("gate.six_path_execution_started False", gate["six_path_execution_started"] is False),
    ("gate.shapley/l4/economics/double_layer False",
     all(gate[k] is False for k in ("shapley_allowed", "l4_allowed", "economics_allowed", "double_layer_uncertainty_allowed"))),
    ("gate.legacy_results_paper_citable False", gate["legacy_results_paper_citable"] is False),
    ("gate.round4_consistency_gate.passed", gate.get("round4_consistency_gate", {}).get("passed") is True),
    ("gate.verified_admitted_count 149", gate["verified_admitted_count"] == 149),
    ("gate.max_frozen_representatives 5", gate["max_frozen_representatives"] == 5),
    ("manifest.cost_sensitivity_groups 80", manifest["cost_sensitivity_groups"] == 80),
    ("manifest.data_points 2016=35136", manifest["data_points_per_year"]["2016"] == 35136),
    ("manifest.tests.functions 127", manifest["tests"]["functions"] == 127),
    ("manifest.tests.cases 136", manifest["tests"]["cases"] == 136),
    ("manifest.tests.stage2b 43/50", manifest["tests"]["breakdown_functions"]["stage2b"] == 43
     and manifest["tests"]["breakdown_cases"]["stage2b"] == 50),
    ("hash_record round4 verification", hashrec.get("round4_verification", {}).get("all_three_hashes_reproducible") is True),
]
for label, ok in json_checks:
    print(f"[{'OK' if ok else 'FAIL'}] {label}: {ok}")
    all_ok &= ok

# cross-doc hash identity
hashes = {k: texts[k].count("11ed37be41737aaba1558b2eed366457d50724c73e9ac7094ea2f9afd249965f")
          for k in ("ledger", "main")}
print("[check] ledger & main both carry governance config hash:", hashes, hashes.get("ledger", 0) >= 1 and hashes.get("main", 0) >= 1)
all_ok &= hashes.get("ledger", 0) >= 1 and hashes.get("main", 0) >= 1

print("\n===== ROUND4 VERIFICATION:", "ALL PASS" if all_ok else "HAS FAILURES", "=====")
raise SystemExit(0 if all_ok else 1)
