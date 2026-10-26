# -*- coding: utf-8 -*-
"""Round-7a pre-update of 统一结果台账.docx (parallel with ④ double-layer run):
1) S8: units unified -> comparable magnitudes (effect_pct across all 6 paths)
2) S12: test count 136 -> 149 (EFC 13)
3) Table3 append R9 record
"""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\统一结果台账.docx")

R9 = [
    "round7a 预更新（2026-08-04，与 ④ 嵌套双层并行）：cross_device_paths.json 已统一六路径单位为 effect_pct（相对目标设备自身退化的百分比贡献，D∈{0.3,0.5,0.7} 三水平）；统一单位下 α₁/β₁ 量级大于跨设备路径（α₁ 最大 +24.33% @ hydrogen_dominant D=0.7；跨设备 |effect_pct|≤~3%）。原“单位未统一、不做主导断言”表述更新为“单位已统一、可比较”",
    "round7a：测试计数 136→149（新增 test_bat_efc_cycle_model.py 13 用例，函数 127→140）；EFC 模型（BAT_CYCLE_MODEL=\"efc\"）已为 closed_loop_lite.py 默认",
    "证据：pytest 全量 149 passed；cross_device_paths.json 6 路径×3 水平×effect_pct；门控字段翻转仍待 ④ 双层（256×4000）完成后统一进行",
]


def main() -> int:
    doc = docx.Document(str(DOC))
    t1 = doc.tables[1]

    for row in t1.rows:
        c0 = row.cells[0].text.strip()
        if c0.startswith("S8"):
            row.cells[5].text = "effect_pct（统一为相对目标设备自身退化的百分比贡献）；D∈{0.3,0.5,0.7} 三水平"
            row.cells[7].text = (
                "六路径单位已统一（effect_pct 百分比贡献，D∈{0.3,0.5,0.7} 三水平）；"
                "统一单位下 α₁/β₁（BAT 源路径）量级大于跨设备路径：α₁ 最大 +24.33% "
                "@ hydrogen_dominant D=0.7、β₁ 最大 +10.47% @ critical D=0.7；"
                "跨设备路径 |effect_pct| ≤ ~3%（α₂ −1.71%~+2.97%，β₂/γ₁/γ₂ <1.5%）；"
                "可比较量级，但不能断言跨设备路径“可忽略”或“仅 2 条激活”"
            )
            print("S8 updated")
        elif c0.startswith("S12"):
            row.cells[4].text = "08_tests/（140 函数 → 149 用例，含 EFC 13）"
            row.cells[6].text = "149 全 PASS"
            row.cells[7].text = (
                "96 基础（Stage2b 50 + unit 34 + Stage3 12）+ 40 专项"
                "（Shapley C/P 6 + 退化时间单位 6 + EOL 系统状态 21 + L4 控制器 7）"
                " + 13 EFC（test_bat_efc_cycle_model.py）；"
                "其中 Stage2b 八文件 = 43 函数/50 用例、unit = 32 函数/34 用例（用例含 e2e 参数化展开）"
            )
            print("S12 updated")

    t3 = doc.tables[3]
    new_row = t3.add_row()
    for ci, txt in enumerate(R9):
        new_row.cells[ci].text = txt
    print("Table3 appended row:", len(t3.rows) - 1)

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
