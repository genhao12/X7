"""docx 文档构建辅助函数。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_background(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=9, color=None, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        sizes = {0: 20, 1: 15, 2: 12.5, 3: 11}
        run.font.size = Pt(sizes.get(level, 11))
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return h


def add_para(doc, text, size=10.5, bold=False, italic=False, color=None, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return p


def add_data_table(doc, rows_data):
    headers = ['数据来源（DOI / 名称 / 作者 / 年份）', '数据内容（曲线 / 数据集 / 模型 / 额定点）',
               '用途（用于哪部分）', '可信度', '可信度说明']
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr_cells[i], h, bold=True, size=9.5,
                      color=RGBColor(0xFF, 0xFF, 0xFF),
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_background(hdr_cells[i], '1F3A5F')
    cred_colors = {'高': 'C6EFCE', '中': 'FFEB9C', '低': 'FFC7CE', '不可用': 'D9D9D9', '缺口': 'D9D9D9'}
    for row_data in rows_data:
        cells = table.add_row().cells
        set_cell_text(cells[0], row_data['source'], size=9)
        set_cell_text(cells[1], row_data['content'], size=9)
        set_cell_text(cells[2], row_data['purpose'], size=9)
        cred = row_data['credibility']
        set_cell_text(cells[3], cred, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        bg = 'D9D9D9'
        for key, val in cred_colors.items():
            if key in cred:
                bg = val
                break
        set_cell_background(cells[3], bg)
        set_cell_text(cells[4], row_data['credibility_reason'], size=9)
    return table


def build_doc(sections, intro_paras, conclusion_paras, out_path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for section in doc.sections:
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    add_heading(doc, '离网 PV–H₂–LFP 微电网项目数据来源与可信度台账', level=0)
    add_para(doc, '版本：2026-07-29  |  适用范围：Phase A 方法与情景研究',
             size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66),
             align=WD_ALIGN_PARAGRAPH.CENTER)

    add_heading(doc, '一、文档说明与可信度评级标准', level=1)
    for p in intro_paras:
        add_para(doc, p['text'], size=p.get('size', 10.5), bold=p.get('bold', False),
                 italic=p.get('italic', False), color=p.get('color'),
                 align=p.get('align'))

    doc.add_page_break()
    for sec in sections:
        add_heading(doc, sec['title'], level=1)
        for sub in sec['sub']:
            add_heading(doc, sub['subtitle'], level=2)
            add_data_table(doc, sub['rows'])
        if sec.get('page_break_after', False):
            doc.add_page_break()

    # 结论
    doc.add_page_break()
    add_heading(doc, '十、当前数据可实现的研究目标与边界', level=1)
    for p in conclusion_paras:
        if p.get('heading'):
            add_heading(doc, p['text'], level=2)
        else:
            add_para(doc, p['text'], size=p.get('size', 10.5), bold=p.get('bold', False),
                     italic=p.get('italic', False), color=p.get('color'))

    doc.save(out_path)
    print(f'文档已生成: {out_path}')
