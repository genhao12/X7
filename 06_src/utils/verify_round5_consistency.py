# -*- coding: utf-8 -*-
"""Round-5 final verification across ledger/summary/main-report/governance JSON."""
from __future__ import annotations

import json
from pathlib import Path

import docx

ROOT = Path(r"d:\matlab\C")
docs = {
    "ledger": ROOT / "09_docs" / "统一结果台账.docx",
    "summary": ROOT / "09_docs" / "工作区运行结果汇总.docx",
    "main": ROOT / "09_docs" / "工作区全面总结报告.docx",
}


def collect(path: Path) -> list[tuple[str, str]]:
    d = docx.Document(str(path))
    out = [(f"P{i}", p.text) for i, p in enumerate(d.paragraphs)]
    for ti, tb in enumerate(d.tables):
        for ri, row in enumerate(tb.rows):
            for ci, c in enumerate(row.cells):
                out.append((f"T{ti}R{ri}C{ci}", c.text))
    return out


texts = {k: "\n".join(t for _, t in collect(p)) for k, p in docs.items()}
all_ok = True


def check(label: str, ok: bool):
    global all_ok
    print(f"[{'OK' if ok else 'FAIL'}] {label}")
    all_ok &= bool(ok)


# --- A. magnitude-leading claims removed (must be absent) ---
absents = {
    "α₁/β₁ 主导；α₂/β₂/γ₁/γ₂ 四条跨设备路径均非零": ["ledger"],
    "主导且量级大": ["summary"],
    "α₁/β₁（ΔD/yr）主导；α₂/β₂/γ₁/γ₂ 四条跨设备路径均非零": ["summary"],
    "β₁(BAT→EL) 主导；α₂/β₂/γ₁/γ₂ 跨设备": ["summary"],
    "为主导路径；跨设备路径": ["main"],
    "α₁/β₁（ΔD/yr）主导；α₂/β₂/γ₁/γ₂为百分比贡献": ["main"],
}
for needle, dks in absents.items():
    for dk in dks:
        check(f"{dk} MUST NOT contain {needle[:25]!r}", needle not in texts[dk])

# --- B. corrected claims present ---
presents = {
    "ledger": ["main_domain_failed=True", "stage2c_admission_passed=True", "不作物量级主导断言", "round5"],
    "summary": ["单位未统一、量级不可比", "不作物量级主导断言"],
    "main": ["不作物量级主导断言", "方法层计算已落地（six_path_execution_started=False", "shapley_allowed=False", "l4_allowed=False", "由未四舍五入的原始值计算", "economics_allowed=l4_allowed=False"],
}
for dk, needles in presents.items():
    for n in needles:
        check(f"{dk} contains {n[:30]!r}", n in texts[dk])

# --- C. governance JSON ---
gov = ROOT / "07_results" / "rectification_v2_stage2c" / "governance"
gate = json.loads((gov / "downstream_gate.json").read_text(encoding="utf-8"))
manifest = json.loads((gov / "run_manifest.json").read_text(encoding="utf-8"))
hashrec = json.loads((gov / "hash_record.json").read_text(encoding="utf-8"))

check("gate.gate_passed False", gate["gate_passed"] is False)
check("gate.upstream_gate.main_domain_failed True", gate["upstream_gate"]["main_domain_failed"] is True)
check("gate.upstream_gate.stage2c_admission_passed True", gate["upstream_gate"]["stage2c_admission_passed"] is True)
check("gate.degradation_execution_started False", gate["degradation_execution_started"] is False)
check("gate.six_path_execution_started False", gate["six_path_execution_started"] is False)
check("gate all execution flags False",
      all(gate[k] is False for k in ("shapley_allowed", "l4_allowed", "economics_allowed",
                                     "double_layer_uncertainty_allowed", "legacy_results_paper_citable")))
check("gate.round5_consistency_gate.passed", gate.get("round5_consistency_gate", {}).get("passed") is True)
check("gate.round5 magnitude_claims_removed", gate.get("round5_consistency_gate", {}).get("magnitude_leading_claims_removed") is True)
check("manifest.upstream_gate split", manifest["upstream_gate"]["main_domain_failed"] is True
      and manifest["upstream_gate"]["stage2c_admission_passed"] is True)
check("manifest facts (80/35136/127/136/43/50)",
      manifest["cost_sensitivity_groups"] == 80
      and manifest["data_points_per_year"]["2016"] == 35136
      and manifest["tests"]["functions"] == 127 and manifest["tests"]["cases"] == 136
      and manifest["tests"]["breakdown_functions"]["stage2b"] == 43
      and manifest["tests"]["breakdown_cases"]["stage2b"] == 50)
check("hash_record round5 reproducible", hashrec.get("round5_verification", {}).get("all_three_hashes_reproducible") is True)
check("contract hashes unchanged",
      gate["input_hash"] == manifest["input_hash"] == hashrec["input_hash"] == "86e12930b6b08cab63dd00e12d5e00ab5afd0a3d1cc6e1122c90b1b52d207dd3"
      and gate["config_hash"] == "11ed37be41737aaba1558b2eed366457d50724c73e9ac7094ea2f9afd249965f"
      and gate["script_hash"] == "42c21ad0a5e6f8451956da552455bf580305709b3ccad85c62c44c00b0b7cb11")

print("\n===== ROUND5 VERIFICATION:", "ALL PASS" if all_ok else "HAS FAILURES", "=====")
raise SystemExit(0 if all_ok else 1)
