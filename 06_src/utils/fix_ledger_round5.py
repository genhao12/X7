# -*- coding: utf-8 -*-
"""Round-5 fixes: 台账 (统一结果台账.docx).
1) T0R10 upstream_gate split into main_domain_failed + stage2c_admission_passed
2) T1R8 (S8) remove "α₁/β₁ 主导" magnitude-leading claim
3) Table3 append R8 round-5 fix record
4) trailer paragraph for round-5 verification
"""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\统一结果台账.docx")

R8 = [
    "round5 前：T0R10 upstream_gate=frozen_capacity_admission_failed 把主域失败与 Stage2c 149 组准入混为一项；T1R8（S8）及主/汇总多处写“α₁/β₁ 主导/量级大”但单位未统一，量级不可比；主报告 Table22/23 L4 百分比由未四舍五入原始值计算，表面 4 位小数值复算不一致",
    "round5 修正：T0R10 拆分 main_domain_failed=True（主域 0 准入，主域下游冻结）+ stage2c_admission_passed=True（149 准入/5 冻结）；全部“α₁/β₁ 主导/量级更大”结论删除，改为“单位未统一、量级不可比、不做主导断言”（台账/汇总/主报告）；Table22/23 加注“变化%由未四舍五入原始值计算，精确值见 l4_ablation.json”；gate/run_manifest/hash_record 重新生成并新增 upstream 拆分字段",
    "治理文件：downstream_gate.json(.4)/run_manifest.json(.3)/hash_record.json(.3)；upstream_gate_main_domain_failed=True、stage2c_admission_passed=True；契约三哈希重算可复现不变",
]

TRAILER = (
    "round5 一致性门控复核（2026-08-04 15:1x）：门控状态不变（gate_passed=False；"
    "stage3_basic_degradation_allowed / robust_six_path_planning_allowed=True 仅计划级，"
    "对应 degradation / six_path_execution_started=False；shapley / l4 / economics / "
    "double_layer_uncertainty / legacy_results_paper_citable=False）。"
    "upstream_gate 已拆分：main_domain_failed=True（主域 0 准入，主域下游物理推断冻结）"
    "＋ stage2c_admission_passed=True（扩展域 149 准入、5 冻结代表）。"
    "“α₁/β₁ 主导/量级更大”结论已全部删除（单位未统一、量级不可比）。"
    "L4 表百分比已注明由未四舍五入原始值计算（精确值以 l4_ablation.json 为准）。"
    "此前仅可进行 BAT-EFC 模型升级的规划与 TDD 用例编写，不得启动结果重算。"
)


def main() -> int:
    doc = docx.Document(str(DOC))

    # 1) T0 R10 upstream_gate split
    t0 = doc.tables[0]
    t0.rows[10].cells[0].text = "上游门控 upstream_gate（拆分）"
    t0.rows[10].cells[1].text = (
        "main_domain_failed=True（Stage0-2 主域 160–320 kWp 0 准入 → 主域下游物理推断冻结）；"
        "stage2c_admission_passed=True（Stage2c 扩展域 350–500 kWp 149 组准入、5 冻结代表，S2）"
    )
    print("T0R10 split done")

    # 2) T1 S8 row (contains "α₁/β₁ 主导；α₂/β₂/γ₁/γ₂")
    t1 = doc.tables[1]
    for row in t1.rows:
        if row.cells[0].text.strip().startswith("S8"):
            old = row.cells[7].text
            row.cells[7].text = (
                "α₁/β₁（ΔD/yr 绝对）与 α₂/β₂/γ₁/γ₂（百分比贡献）单位未统一、量级不可比，不作物量级主导断言；"
                "六路径均非零（数据含 D∈{0.3,0.5,0.7} 三水平，effect_pct 多数 <1%、个别达 +2.97% "
                "@ hydrogen_dominant D=0.7）；不能断言跨设备路径“可忽略”或“仅 2 条激活”，需待统一后比较"
            )
            print("S8 remark rewritten; old:", old[:30])
            break

    # 3) Table3 append R8
    t3 = doc.tables[3]
    new_row = t3.add_row()
    for ci, txt in enumerate(R8):
        new_row.cells[ci].text = txt
    print("Table3 appended row:", len(t3.rows) - 1)

    # 4) trailer
    doc.add_paragraph(TRAILER)
    print("trailer appended")

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
