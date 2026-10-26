# -*- coding: utf-8 -*-
"""Round-7a pre-update of 工作区运行结果汇总.docx (parallel with ④):
1) test count 136 -> 149 (EFC 13; 127 fn -> 140 fn)
2) units unified statements (P47, Table5 R3, Table12 R4)
3) Table13 append EFC row, total 149
"""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\工作区运行结果汇总.docx")


def set_para(doc, needle, new_text):
    for para in doc.paragraphs:
        if needle in para.text:
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
            print(f"[PARA] {needle[:24]!r} replaced")
            return True
    print(f"[WARN] para not found: {needle!r}")
    return False


def main() -> int:
    doc = docx.Document(str(DOC))

    # 1) P5 test count
    set_para(doc, "全量测试 136 项 PASS（127 个测试函数经参数化展开）",
             "全局结论：工作区现有运行结果全部可复现且关键闭环验证通过。全量测试 149 项 PASS"
             "（140 个测试函数经参数化展开，含 EFC 13 项）；跨阶段验证 46 项 PASS、跨模块验证 55 项 PASS、"
             "VS1–VS8 系列共 33 项检查 PASS。所有经济/寿命/耦合结论均来自单一真值源脚本"
             "（economics.npc_lite.compute_equipment_capex、Stage3 EOL 事件模型、EFC 循环损伤模型）。")

    # 2) P42 test decomposition
    set_para(doc, "说明：报告早期曾误写",
             "说明：报告早期曾误写「50+34+12=136」（实际=96），正确口径为 Stage2b 50 + unit 34 + Stage3 12"
             " = 96 基础项，加 4 个专项文件共 40 项，再加 EFC 循环损伤 13 项 = 149。参数化（@parametrize）"
             "将 140 个函数展开为 149 个执行用例。")

    # 3) P47 units unified
    p47 = (
        "六路径单位已统一（effect_pct，相对目标设备自身退化的百分比贡献，D∈{0.3,0.5,0.7} 三水平）；"
        "统一单位下 α₁/β₁（BAT 源路径）量级大于跨设备路径：α₁ 最大 +24.33% @ hydrogen_dominant D=0.7、"
        "β₁ 最大 +10.47% @ critical D=0.7；跨设备路径 |effect_pct| ≤ ~3%（α₂ −1.71%~+2.97%，"
        "β₂/γ₁/γ₂ <1.5%）；可比较量级，但不能断言跨设备路径“可忽略”或“仅 2 条激活”。"
    )
    set_para(doc, "六路径中 α₁/β₁（BAT→FC/EL，ΔD/yr 绝对）", p47)

    # 4) Table5 VS3 row C2/C3
    for tb in doc.tables:
        if len(tb.rows) > 0 and tb.rows[0].cells[0].text.strip() == "验证项":
            for row in tb.rows:
                if row.cells[0].text.strip() == "VS3 路径识别":
                    row.cells[2].text = (
                        "α₁/β₁ 与 α₂/β₂/γ₁/γ₂ 单位已统一（effect_pct 百分比贡献）；"
                        "α₁/β₁ 量级大于跨设备路径（α₁ 最大 +24.33%）；跨设备路径均非零、|effect_pct|≤~3%；"
                        "C/P 归因仅 4 联盟；D∈{0.3,0.5,0.7} 三水平"
                    )
                    row.cells[3].text = (
                        "统一单位下 α₁/β₁ 量级大于跨设备路径；跨设备路径均非零（|effect_pct|≤~3%，"
                        "α₂ −1.71%~+2.97%）；不能断言跨设备“可忽略”或“仅 2 条激活”"
                    )
                    print("VS3 row updated")
            break

    # 5) Table12 R4
    for tb in doc.tables:
        if len(tb.rows) > 1 and tb.rows[1].cells[0].text.strip() == "脚本":
            r1 = tb.rows[1].cells[1].text
            if "run_cross_device_paths" in r1:
                for row in tb.rows:
                    if row.cells[0].text.strip() == "主要结果":
                        row.cells[1].text = (
                            "六路径单位已统一（effect_pct 百分比贡献，D∈{0.3,0.5,0.7} 三水平）："
                            "α₁(BAT→FC)、β₁(BAT→EL) 量级大于跨设备路径（α₁ 最大 +24.33% @ hydrogen_dominant D=0.7；"
                            "β₁ 最大 +10.47% @ critical D=0.7）；α₂/β₂/γ₁/γ₂ 均非零（α₂ −1.71%~+2.97%，"
                            "β₂/γ₁/γ₂ <1.5%）；可比较量级，不能断言跨设备“可忽略”或“仅 2 条激活”"
                        )
                        print("Table12 R4 updated")
                break

    # 6) Table13 append EFC row + total
    for tb in doc.tables:
        if len(tb.rows) > 0 and tb.rows[0].cells[0].text.strip() == "分组":
            # append EFC row before total row
            total_row = None
            for row in tb.rows:
                if row.cells[0].text.strip() == "合计":
                    total_row = row
                    break
            if total_row is not None:
                # insert new row before total by creating cells
                import copy
                new_tr = copy.deepcopy(total_row._tr)
                total_row._tr.addprevious(new_tr)
                from docx.table import _Row
                nr = _Row(new_tr, tb)
                vals = ["EFC 循环损伤（1 文件）", "test_bat_efc_cycle_model(13)（rainflow+EFC，280Ah 校准）", "13", "通过"]
                for ci, v in enumerate(vals):
                    nr.cells[ci].text = v
                # update total row
                total_row.cells[1].text = "96 + 40 + 13"
                total_row.cells[2].text = "149"
                total_row.cells[3].text = "全部 PASS（140 函数经 @parametrize 展开为 149 用例）"
                print("Table13 EFC row inserted, total updated")
            break

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
