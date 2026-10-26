# -*- coding: utf-8 -*-
"""Round-5 fixes: 汇总报告 (工作区运行结果汇总.docx).
Remove "α₁/β₁ 主导/量级大" magnitude-leading claims where units are not unified.
"""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\工作区运行结果汇总.docx")


def main() -> int:
    doc = docx.Document(str(DOC))

    # 1) P47 paragraph
    for para in doc.paragraphs:
        if "六路径中 BAT→FC/EL" in para.text:
            new_p47 = (
                "六路径中 α₁/β₁（BAT→FC/EL，ΔD/yr 绝对）与四条跨设备反向路径"
                "（α₂/β₂/γ₁/γ₂，百分比贡献）单位未统一、量级不可比，不作物量级主导断言；"
                "跨设备路径均非零（多数 |effect_pct|<1%、个别达 +2.97% @ hydrogen_dominant D=0.7）；"
                "不能断言跨设备路径“可忽略”或“仅 2 条激活”，需待统一后比较。"
            )
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_p47
            print("P47 rewritten")
            break

    # 2) Table5 R3 result cell (VS3)
    for tb in doc.tables:
        if len(tb.rows) > 0 and tb.rows[0].cells[0].text.strip() == "验证项":
            for row in tb.rows:
                if row.cells[0].text.strip() == "VS3 路径识别":
                    row.cells[3].text = (
                        "α₁/β₁（ΔD/yr）与 α₂/β₂/γ₁/γ₂（百分比贡献）单位未统一、量级不可比；"
                        "跨设备路径均非零；C/P 归因仅 4 联盟；仅展示 D=0.5，不作物量级主导断言"
                    )
                    print("VS3 R3 result rewritten")
            break

    # 3) Table12 R4 (六路径主要结果)
    for tb in doc.tables:
        if len(tb.rows) > 0 and tb.rows[1].cells[0].text.strip() == "脚本":
            r1 = tb.rows[1].cells[1].text
            if "run_cross_device_paths" in r1:
                for row in tb.rows:
                    if row.cells[0].text.strip() == "主要结果":
                        row.cells[1].text = (
                            "α₁(BAT→FC)、β₁(BAT→EL) 为 ΔD/yr 绝对单位，α₂/β₂/γ₁/γ₂ 跨设备路径为"
                            "百分比贡献（均非零，effect_pct 多数 <1%、个别 +2.97%；数据含 "
                            "D∈{0.3,0.5,0.7} 三水平）；单位未统一、量级不可比，不作物量级主导断言，"
                            "不能断言“可忽略”或“仅 2 条激活”，需待统一后比较"
                        )
                        print("Table12 R4 rewritten")
                break

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
