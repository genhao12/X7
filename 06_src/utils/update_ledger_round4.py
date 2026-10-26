# -*- coding: utf-8 -*-
"""Round-4 update of 统一结果台账.docx:
1) Table2 R10 gate_reason synced with regenerated downstream_gate.json
2) Table3 appends R7 (round-4 fix record)
3) trailing paragraph documenting round-4 consistency verification
"""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\统一结果台账.docx")

NEW_GATE_REASON = (
    "capacity_admission_passed_stage3_plan_only；round4 单一真值源一致性门控 PASS"
    "（台账/汇总/主报告对齐：2016=35136 闰年、5×16=80 组、127 函数→136 用例、"
    "Stage2b 43 函数/50 用例、六路径非零且单位未统一、已移除“可忽略/仅2条激活”断言）；"
    "所有执行/论文标志仍为 False"
)

R7 = [
    "主报告 round4 前：P61 测试数仅写“50项”未区分 43 函数/50 用例；P104 标题“仅α₁/β₁”易误读为仅 2 条激活；P125 “γ₁/γ₂ 可忽略（<0.1%）”无 D=0.5 限定；Table24 状态“完成（仅α₁/β₁）”；Table25 哈希为单文件口径（a17188be… 等）",
    "主报告已修正（2026-08-04 round4）：127 函数→136 用例、43/50、32/34 分列；10.4 标题注明跨设备见 10.4b；γ₁/γ₂ 加 D_source=0.5 限定并声明全水平待统一后补证；Table24 改“六路径已执行”并标注 shapley_allowed=False / l4_allowed=False；Table25 哈希统一治理口径（与 Table0 一致）；汇总报告 Table9/11/12“是否通过”补“方法层”前缀",
    "fix_main_report_round4.py / fix_summary_report_round4.py；治理文件重新生成：downstream_gate.json(.3) / run_manifest.json(.2) / hash_record.json(.2)，契约三哈希重算可复现不变（input=86e12930… config=11ed37be… script=42c21ad0…）",
]

TRAILER = (
    "round4 一致性核验（2026-08-04 14:50）：台账、运行结果汇总、全面总结报告三份文档已对齐；"
    "downstream_gate.json / run_manifest.json / hash_record.json 已重新生成，"
    "契约三哈希重算可复现不变；门控仍未通过（所有执行/论文标志 False），"
    "允许进入的仅为规划级（stage3_basic_degradation_allowed / robust_six_path_planning_allowed=True，"
    "对应 execution_started=False），BAT-EFC 升级须待门控通过后启动。"
)


def main() -> int:
    doc = docx.Document(str(DOC))

    # 1) Table2 gate_reason row
    t2 = doc.tables[2]
    for row in t2.rows:
        if row.cells[0].text.strip() == "gate_reason":
            old = row.cells[1].text
            row.cells[1].text = NEW_GATE_REASON
            print("Table2 gate_reason updated:", old[:30], "->", NEW_GATE_REASON[:30])
            break

    # 2) Table3 append R7
    t3 = doc.tables[3]
    new_row = t3.add_row()
    for ci, txt in enumerate(R7):
        new_row.cells[ci].text = txt
    print("Table3 appended row:", len(t3.rows) - 1)

    # 3) trailer paragraph at end of document body
    doc.add_paragraph(TRAILER)
    print("trailer appended")

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
