# -*- coding: utf-8 -*-
"""生成整改阶段结果与未决问题报告 Word 文档"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)
style = doc.styles['Normal']; style.font.name = 'Times New Roman'; style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6); style.paragraph_format.line_spacing = 1.15
for i in range(1, 4):
    hs = doc.styles[f'Heading {i}']; hs.font.name = 'Times New Roman'
    hs.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers), style='Light Grid Accent 1')
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
        for p in table.rows[0].cells[j].paragraphs:
            for r in p.runs: r.font.bold = True; r.font.size = Pt(9)
    for row_data in rows:
        row = table.add_row()
        for j, val in enumerate(row_data):
            row.cells[j].text = str(val)
            for p in row.cells[j].paragraphs:
                for r in p.runs: r.font.size = Pt(9)

# 封面
title = doc.add_heading('', level=0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('离网光伏-氢-电池混合储能微电网\n整改阶段结果与未决问题报告')
run.font.size = Pt(20); run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('文档性质：整改阶段结果，方法框架保留但未完成有效验证\n生成日期：2026-08-02').font.size = Pt(12)
doc.add_page_break()

# 第一章
doc.add_heading('第一章 最终判定', level=1)
doc.add_paragraph(
    '方法框架可以保留，但以下5项均未完成有效验证：'
    '容量闭合、FC/BAT退化标定、六路径统计判据、Shapley算术闭合和经济事件模型。'
    '当前文档应定位为"整改阶段结果与未决问题报告"。'
)
doc.add_heading('1.1 已正确修正的方法骨架', level=2)
correct = [
    '路径符号统一(v3.0冻结定义)',
    '固定D路径辨识与动态寿命评价协议分离',
    'EOL增量封顶(已修正为逐时间步执行)',
    '设备级损伤表达(TE向量替代总损伤求和)',
]
for c in correct:
    doc.add_paragraph(c, style='List Bullet')

doc.add_heading('1.2 未完成有效验证的5项', level=2)
add_table(
    ['项目', '问题', '修正要求'],
    [
        ['容量闭合', 'H2年末894kg<年初1000kg, LPSP依赖初始库存补贴', '实现周期稳态, 排除外部能量补贴'],
        ['FC/BAT退化标定', '999h数据换算缺观测变量/拟合误差/验证集, ±20%无推导', '只能称情景标定, 不能升级B级'],
        ['六路径统计判据', 'AE=0只证明代码输出零, 缺NC1/NC3/Bootstrap/阈值/方向一致率', '补全7项激活标准'],
        ['Shapley算术闭合', 'phi_C+phi_P=0.004453≠alpha1的0.004738, 相差6%', '明确两表对应关系或修正'],
        ['经济事件模型', 'FC年运行6462h/EOL=1278h→年5次更换, 但模型只算20次', '修正时间推进逻辑'],
    ]
)
doc.add_page_break()

# 第二章
doc.add_heading('第二章 容量闭合未决问题', level=1)
doc.add_heading('2.1 H2容量/初值混用', level=2)
doc.add_paragraph(
    'C6配置储氢容量500kg，但闭合验证报告H2库存范围608–1000kg——直接超过上限500kg。'
    '此矛盾源于闭合验证脚本使用了C5配置(H2=1000kg)而非C6配置(H2=500kg)。'
    '需修正：所有验证脚本必须使用一致的配置参数。'
)
doc.add_heading('2.2 年度能量未闭合', level=2)
doc.add_paragraph(
    'H2年末894kg < 年初1000kg（C5配置），系统消耗初始储氢10.6%作为能量补贴。'
    'LPSP=0.24%不能证明系统年度能量闭合，而是依赖初始库存补贴。'
    '修正要求：实现周期稳态（H2年末=年初或差值可由PV盈余解释），排除外部能量补贴。'
)
doc.add_heading('2.3 LPSP数值不一致', level=2)
add_table(
    ['来源', 'C6 LPSP', '说明'],
    [
        ['S7配置表', '0.42%', '搜索脚本结果'],
        ['S3闭合验证', '0.24%', '使用C5配置(非C6)'],
        ['S11跨年', '0.43%', '2017年结果'],
        ['扩展1 L4', '0.43%', '2017年M2结果'],
        ['2016年', '3.09-3.17%', '超1%约束'],
    ]
)
doc.add_paragraph('C6/C7在2016年LPSP>1%，"C6/C7/C5均验证LPSP≤1%"是错误表述。')
doc.add_heading('2.4 C6"接近原目标等级"不准确', level=2)
doc.add_paragraph(
    'C6的FC/EL 80kW和H2 500kg分别约为原上限(45kW/40kg)的1.8倍和12.5倍。'
    '"接近原目标等级"不准确，应改为"相对C5更紧凑但仍超出原搜索域"。'
)
doc.add_page_break()

# 第三章
doc.add_heading('第三章 FC/BAT退化标定未决问题', level=1)
doc.add_heading('3.1 FC标定缺陷', level=2)
add_table(
    ['问题', '说明', '修正要求'],
    [
        ['观测变量不明', '999h数据换算成68.53%/年"效率退化率"，但未说明观测的是电压/效率/功率', '明确观测变量'],
        ['拟合误差缺失', '无拟合R²、RMSE、验证集', '补充拟合误差和验证集'],
        ['参数后验缺失', '无MCMC/Bootstrap后验', '补充参数不确定性'],
        ['系统尺度验证缺失', '999h数据到100kW的迁移无独立验证', '只能称情景标定'],
        ['10%EOL不成立', '10%电压EOL≠10%效率损失', '分别定义电压EOL和效率EOL'],
        ['±20%无推导', '不确定性区间无推导依据', '给出推导或标注为假设'],
    ]
)
doc.add_paragraph('结论：FC退化只能称"文献数据约束下的情景标定"，不能升级为已验证B级结论。')

doc.add_heading('3.2 BAT退化异常', level=2)
doc.add_paragraph(
    'BAT年度损伤0.769、20年更换15次——说明pack换算尚未通过。'
    '280Ah电芯级数据迁移到系统级需核验热梯度/并联不一致/BMS/变换器影响。'
    'BAT迁移后证据应标B级（非A级）。'
)
doc.add_page_break()

# 第四章
doc.add_heading('第四章 经济事件模型未决问题', level=1)
doc.add_heading('4.1 FC更换次数计算错误', level=2)
doc.add_paragraph(
    'FC年运行6462h，EOL=1278h，理论上一年触发约5次更换(6462/1278≈5.06)。'
    '但经济模型只计算20年20次FC更换（每年1次），明显采用了"每年最多更换一次"的错误时间推进。'
    '修正要求：按实际运行小时/EOL小时计算更换次数，或实现EOL后停机/更换事件。'
)
doc.add_heading('4.2 CAPEX分项矛盾', level=2)
add_table(
    ['分项', '金额(万元)', '说明'],
    [
        ['PV', '167.5', '200kWp×8375元/kWp'],
        ['BAT', '108.8', '500kWh×2175元/kWh'],
        ['FC', '116.0', '80kW×14500元/kW'],
        ['EL', '116.0', '80kW×14500元/kW'],
        ['H2', '36.2', '500kg×725元/kg'],
        ['设备小计', '544.5', '—'],
        ['电力电子(15%)', '81.7', '基于设备小计'],
        ['安装(20%)', '108.9', '基于设备小计'],
        ['总CAPEX', '735.1', '—'],
    ]
)
doc.add_paragraph('设备小计544.5万+PE 81.7万+安装108.9万=735.1万，分项合计正确。但H2 500kg只计36万元，未明确储罐/压缩/干燥等工程边界。')
doc.add_heading('4.3 M2/M3 NPC矛盾', level=2)
doc.add_paragraph('M2与M3声称NPC完全相同，却又相差一次EL更换(12 vs 11次)。更换次数不同则NPC不应相同。需修正。')
doc.add_page_break()

# 第五章
doc.add_heading('第五章 六路径统计判据未决问题', level=1)
doc.add_heading('5.1 AE=0不充分', level=2)
doc.add_paragraph(
    '精确显示0.000000只证明当前代码输出为零，不证明路径不存在。'
    'EMS"读取参数"不能排除控制规则未实际响应该参数。'
    '四条路径仍可能是结构性零或数值零。'
)
doc.add_heading('5.2 缺失的7项激活标准', level=2)
add_table(
    ['标准', '说明', '当前状态'],
    [
        ['绝对效应阈值', 'AE需超过模型不确定性', '缺失'],
        ['95%CI不跨0', 'Bootstrap区间', '缺失'],
        ['方向一致率≥90%', '跨窗口/跨年', '缺失'],
        ['NC1负对照', '联盟OFF残差', '缺失'],
        ['NC3负对照', '归因冻结验证', '缺失'],
        ['归因闭合误差≤5%', 'Shapley闭合', '部分(残差=0但算术不一致)'],
        ['交互占比≤30%', '二阶交互', '缺失'],
    ]
)
doc.add_heading('5.3 α₁/β₁的AE需证明高于不确定性', level=2)
doc.add_paragraph(
    'α₁的AE=0.004738，β₁的AE=0.003329，必须证明其高于参数及模型不确定性。'
    '当前无Bootstrap区间、无负对照、无跨配置复现，不能确认为激活。'
)
doc.add_page_break()

# 第六章
doc.add_heading('第六章 Shapley算术闭合未决问题', level=1)
doc.add_heading('6.1 归因变量定义冲突', level=2)
add_table(
    ['来源', '归因定义', '问题'],
    [
        ['第3页六路径', '母线纹波/供氢波动', '物理归因'],
        ['第6页Shapley', '容量C/效率E/功率P/库存I', '计算归因'],
        ['冲突', '两套物理定义不对应', '需明确对应关系'],
    ]
)
doc.add_heading('6.2 算术不一致', level=2)
doc.add_paragraph(
    'Shapley表中 phi_C + phi_P = 0.004350 + 0.000103 = 0.004453，'
    '但α₁总效应为0.004738，相差0.000285（约占6.0%）。'
    '与"闭合残差为0、交互占比低于0.1%"矛盾。'
    '需明确：Shapley归因(C/E/P/I)与六路径(α₁/β₁/γ₁等)是否对应同一输出，或修正算术。'
)
doc.add_page_break()

# 第七章
doc.add_heading('第七章 L4与跨年未决问题', level=1)
doc.add_heading('7.1 L4在C6上验证失败', level=2)
doc.add_paragraph(
    'L4在C6配置下同时增加退化(M3>M2)和LPSP(M3>M2)，应结论为"该控制器在C6上验证失败并被否决"，'
    '不能仅写"配置依赖"。'
)
doc.add_heading('7.2 跨年结果标注缺失', level=2)
doc.add_paragraph(
    '跨年表中的TE_BAT、TE_EL未标明源设备、源健康档位和反事实基准，不能对应到六条有向路径。'
    'C6/C7在2016年本身不满足可靠性约束(LPSP>1%)，不能用于宣称可行系统的跨年稳健性。'
)
doc.add_page_break()

# 第八章
doc.add_heading('第八章 回归测试遗漏', level=1)
doc.add_paragraph(
    '回归测试只显示α₁、α₂、β₁、γ₁四个符号，遗漏β₂和γ₂，不能称六路径符号全部通过。'
    '需补充β₂(FC→EL)和γ₂(EL→BAT)的符号验证。'
)
add_table(
    ['路径', '期望(v3.0)', 'S12测试', '状态'],
    [
        ['α₁', 'BAT→FC', '已测PASS', '✅'],
        ['α₂', 'EL→FC', '已测PASS', '✅'],
        ['β₁', 'BAT→EL', '已测PASS', '✅'],
        ['β₂', 'FC→EL', '未测', '❌遗漏'],
        ['γ₁', 'FC→BAT', '已测PASS', '✅'],
        ['γ₂', 'EL→BAT', '未测', '❌遗漏'],
    ]
)
doc.add_page_break()

# 第九章
doc.add_heading('第九章 EOL守卫修正记录', level=1)
doc.add_paragraph(
    '审查发现EOL守卫写成对672步总损伤一次封顶，无法获得准确EOL时刻和更换事件。'
    '已修正为逐时间步执行：'
)
doc.add_paragraph('dD_used(t) = min[dD_raw(t), 1.0 - D(t)]', style='Intense Quote')
doc.add_paragraph(
    '每步更新D(t+1) = D(t) + dD_used(t)，准确捕获EOL时刻。'
    '此修正使后续经济评价的更换事件计算可基于准确EOL时间。'
)
doc.add_page_break()

# 第十章
doc.add_heading('第十章 下一步修正顺序', level=1)
add_table(
    ['步骤', '任务', '前置条件'],
    [
        ['F1', '修正H2容量/初值混用, 实现周期稳态', '排除外部能量补贴'],
        ['F2', '补充FC标定(观测变量/拟合误差/验证集/后验)', 'F1'],
        ['F3', 'BAT pack换算核验', 'F1'],
        ['F4', '逐步EOL与更换事件实现', 'F2+F3'],
        ['F5', '经济模型修正(FC更换次数/时间推进)', 'F4'],
        ['F6', '六路径补全7项激活标准(NC1/NC3/Bootstrap/阈值)', 'F1-F4'],
        ['F7', 'Shapley归因变量统一(C/E/P/I vs α/β/γ对应)', 'F6'],
        ['F8', '跨年验证(仅用满足LPSP≤1%的配置)', 'F1'],
        ['扩展', 'L4(已否决)/经济(待F5修正)', 'F1-F8'],
    ]
)
doc.add_heading('10.1 当前文档定位', level=2)
doc.add_paragraph(
    '本报告为"整改阶段结果与未决问题报告"。方法框架可保留，但容量闭合、FC/BAT退化标定、'
    '六路径统计判据、Shapley算术闭合和经济事件模型均未完成有效验证。'
    '经济与L4继续作为扩展结果，不作为核心方法论文依据。'
)

doc.add_heading('10.2 F1周期稳态验证结果', level=2)
doc.add_paragraph(
    'F1验证发现根本问题：UCSD 2016/2018年负荷与PV不匹配，系统在非2017年份无法能量闭合。'
)
add_table(
    ['配置', '年份', 'LPSP(%)', 'H2变化(%)', '周期稳态', '说明'],
    [
        ['C6_compact(H2=500)', '2016', '4.01', '-46.1', 'FAIL', 'H2严重消耗'],
        ['C6_compact(H2=500)', '2017', '0.47', '-22.5', 'FAIL', 'H2消耗22%'],
        ['C6_compact(H2=500)', '2018', '0.37', '-27.9', 'FAIL', 'H2消耗28%'],
        ['C5-H2-500(PV300/BAT800)', '2016', '3.65', '-25.6', 'FAIL', '2016年数据不匹配'],
        ['C5-H2-500(PV300/BAT800)', '2017', '0.17', '-2.7', 'PASS', '唯一稳态'],
        ['C5-H2-500(PV300/BAT800)', '2018', '0.14', '-10.3', 'FAIL', 'H2消耗10%'],
    ]
)
doc.add_paragraph(
    '结论：仅2017年C5-H2-500配置达到周期稳态。2016/2018年能量赤字是UCSD数据特性'
    '（负荷高/PV低），非配置问题。后续验证应：'
    '(1)以2017年为基准年；'
    '(2)2016/2018年仅作敏感性参考，不作为稳态验证；'
    '(3)或使用多年滚动仿真验证跨年闭合。'
)
doc.add_paragraph(
    'C6_compact（BAT=500/FC=80/EL=80）在所有年份均能量赤字，该配置不可行。'
    'C5-H2-500（BAT=800/FC=100/EL=100/H2=500）是2017年唯一满足LPSP≤1%+周期稳态的最小配置。'
)

out_path = Path(r'D:\matlab\C\09_docs\工作区全面总结报告.docx')
doc.save(str(out_path))
print(f'文档已生成: {out_path}')
print(f'段落数: {len(doc.paragraphs)}, 表格数: {len(doc.tables)}')
