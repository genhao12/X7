# -*- coding: utf-8 -*-
"""Round-10 final fixes:
1) main report P90 -> full EFC economics paragraph
2) summary Table7 (M4 legacy economics) -> scenario/legacy annotation
3) ledger Table0 append extended-hash row; Table3 append R11 record
"""
from __future__ import annotations

from pathlib import Path

import docx

ROOT = Path(r"d:\matlab\C")


def set_para(doc, needle, new_text):
    for para in doc.paragraphs:
        if needle in para.text:
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
            print(f"[PARA] {needle[:24]!r} replaced")
            return True
    print(f"[WARN] para not found: {needle!r}")
    return False


def main() -> int:
    # ---- main report P90 ----
    mp = ROOT / "09_docs" / "工作区全面总结报告.docx"
    d = docx.Document(str(mp))
    set_para(d, "5配置×16情景=80组下（economics_allowed=True，门控已通过）",
             "经济性基于 EFC 循环损伤模型重算（degradation_summary.json economics，2026-08-04 16:26）："
             "5 配置 × 16 情景 = 80 组（economics_allowed=True，门控已通过）。基线 NPC 31.5~39.3 MCNY、"
             "LCOE 5.55~6.90 CNY/kWh（compact 最低 5.55、hydrogen_dominant 最高 6.90，见 Table14）；"
             "PEMEL 单价为 4/5 配置第一敏感项（per +30% → LCOE +8.5~12.4%），BOP 剥离降 LCOE 19.9~21.2%，"
             "全设备同步缩放 NPC 严格线性。经济结论限定为情景评价（成本参数边界未补全，见下）。")
    d.save(str(mp))
    print("main P90 fixed")

    # ---- summary Table7 legacy economics annotation ----
    sp = ROOT / "09_docs" / "工作区运行结果汇总.docx"
    d2 = docx.Document(str(sp))
    for tb in d2.tables:
        if len(tb.rows) > 1 and "npc_lite" in tb.rows[1].cells[1].text:
            for row in tb.rows:
                label = row.cells[0].text.strip()
                if label == "主要结果":
                    row.cells[1].text = (
                        "初始 CAPEX 319 万元；离散法 NPC=22.21 百万、LCOE=3.944 CNY/kWh；替换成本占 NPC 81.2%"
                        "（legacy MVP 单配置情景值；Stage3 EFC 经济见 degradation_summary.json："
                        "NPC 31.5~39.3M、LCOE 5.55~6.90）"
                    )
                elif label == "是否通过":
                    row.cells[1].text = (
                        "通过（L01–L06：NPC>0、区间合理、CAPEX+OPEX+Repl=NPC、替换>50%；"
                        "legacy MVP 结果，legacy_results_paper_citable=False，仅作方法与情景参考）"
                    )
            print("summary Table7 annotated")
            break
    d2.save(str(sp))
    print("summary saved")

    # ---- ledger Table0 extended hash row + Table3 R11 ----
    lp = ROOT / "09_docs" / "统一结果台账.docx"
    d3 = docx.Document(str(lp))
    t0 = d3.tables[0]
    nr = t0.add_row()
    nr.cells[0].text = "扩展哈希 extended_hashes"
    nr.cells[1].text = ("hash_record.json(.5)：18 个文件 SHA-256 覆盖 EFC 代码（efc_cycle_model/closed_loop_lite）、"
                        "Stage2c 配置（rectification_v2_stage2c.json）、退化/经济参数（frozen/economic/mvp_c1）、"
                        "全部分析脚本（run_stage3_*/run_cost_sensitivity/run_cross_device_paths/run_degradation_uncertainty/"
                        "run_double_layer_propagation/run_l4_ablation/run_stage3_paper_assets/run_efc_vs_proxy_comparison）"
                        "+ ems_stage2b/annual_input_contract/64号CSV；详见 governance/hash_record.json")
    print("ledger Table0 extended-hash row added")

    t3 = d3.tables[3]
    r11 = [
        "round10 修正：S8 联盟数统一为“C/P 仅 4 联盟”（删除 2^4=16 子集歧义）；S9 配置域改为“Stage2c 5 代表配置、真实嵌套 256×4000”（去掉 C1/Stage2b 原型，双层对象确为 5 个 Stage2c 配置）",
        "round10：主报告 Table11/14/16/17 用 degradation_summary.json（EFC，mtime 16:26）与 uncertainty_propagation.json（EFC 参数）刷新寿命/NPC/LCOE/LHS（BAT 寿命 7.2~8.8y、NPC 31.5~39.3M、LCOE 5.55~6.90；旧值 14.3y/29M/5.13 已清除）；汇总 Table7 标注 legacy 情景；门控文档正文全部改已通过；hash_record.json(.5) 扩展 18 文件哈希",
        "证据：degradation_summary.json mtime 16:26 含 economics；uncertainty_propagation.json mtime 16:14 含 BAT_EFC_A_D/K 参数；hash_record.json(.5) extended_hashes 18 项",
    ]
    nr2 = t3.add_row()
    for ci, txt in enumerate(r11):
        nr2.cells[ci].text = txt
    print("ledger Table3 R11 added")

    d3.save(str(lp))
    print("ledger saved")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
