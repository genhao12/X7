# -*- coding: utf-8 -*-
"""Round-4 consistency fixes for 工作区全面总结报告.docx.

Single-truth-source gate: align test counts (43 fn/50 cases, 127 fn/136 cases),
remove "仅α₁/β₁" ambiguity, remove unproven "可忽略" claim, unify hashes to
governance contract values (identical to 统一结果台账 Table0).
"""
from __future__ import annotations

import shutil
from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\工作区全面总结报告.docx")
BACKUP = Path(r"d:\matlab\C\09_docs\工作区全面总结报告_backup_20260804_round4.docx")

GOV_CONFIG_HASH = "11ed37be41737aaba1558b2eed366457d50724c73e9ac7094ea2f9afd249965f"
GOV_SCRIPT_HASH = "42c21ad0a5e6f8451956da552455bf580305709b3ccad85c62c44c00b0b7cb11"


def replace_para_text(doc: docx.Document, match_sub: str, new_text: str) -> bool:
    """Replace a paragraph whose text contains match_sub with new_text."""
    for para in doc.paragraphs:
        if match_sub in para.text:
            # preserve paragraph style, drop run-level formatting (plain text)
            for run in para.runs:
                run.text = ""
            if para.runs:
                para.runs[0].text = new_text
            else:
                para.add_run(new_text)
            print(f"  [PARA] replaced: {match_sub[:30]!r}")
            return True
    print(f"  [WARN] paragraph not found: {match_sub!r}")
    return False


def replace_cell_text(doc: docx.Document, table_startswith: str, row_contains: str,
                      col_new: dict[int, str]) -> bool:
    """Within the first table whose R0 first cell starts with table_startswith,
    find the row whose cell0 contains row_contains, then set cells by index."""
    for tb in doc.tables:
        if not tb.rows:
            continue
        head = tb.rows[0].cells[0].text
        if not head.startswith(table_startswith):
            continue
        for row in tb.rows:
            if row_contains in row.cells[0].text:
                for ci, txt in col_new.items():
                    if ci < len(row.cells):
                        row.cells[ci].text = txt
                print(f"  [CELL] table {table_startswith!r} row {row_contains!r} -> {col_new}")
                return True
        print(f"  [WARN] row not found in table {table_startswith!r}: {row_contains!r}")
        return False
    print(f"  [WARN] table not found: {table_startswith!r}")
    return False


def main() -> int:
    if not BACKUP.exists():
        shutil.copy2(DOC, BACKUP)
    doc = docx.Document(str(DOC))

    print("== P61 test composition ==")
    p61_new = (
        "测试构成（127 个测试函数经 @parametrize 展开为 136 个执行用例，全部通过）："
        "Stage2b（8个文件）43 函数/50 用例 + unit/（12个文件）32 函数/34 用例 + "
        "Stage-3（1个文件）12 用例 = 96 用例基础；另有4个专项文件共 40 用例"
        "（Shapley C/P归因 6 + 退化时间单位 6 + EOL系统状态 21 + L4控制器 7）；"
        "合计 96 + 40 = 136 用例。容量编排器按候选写入原子JSON检查点，"
        "输入哈希、代码哈希或配置哈希任一变化均使检查点失效。"
    )
    replace_para_text(doc, "测试构成（136项全通过）", p61_new)

    print("== P104 heading ==")
    replace_para_text(
        doc,
        "10.4 闭环反事实Shapley分解",
        "10.4 闭环反事实Shapley分解（C/P归因，4联盟；本节覆盖 α₁/β₁，跨设备路径见 10.4b）",
    )

    print("== P125 gamma claim ==")
    p125_new = (
        "D_source=0.5 下 γ₁/γ₂ 百分比贡献 <0.1%（相对目标设备自身退化），"
        "FC/EL 退化对 BAT 损伤影响很小；但正文仅展示 D=0.5，且跨设备路径与 α₁/β₁ 单位未统一，"
        "全水平（0.3/0.5/0.7）的“可忽略”结论须统一单位后补证"
    )
    replace_para_text(doc, "γ₁和γ₂效应可忽略", p125_new)

    print("== Table24 R4/R5 ==")
    replace_cell_text(
        doc, "分析项",
        "闭环Shapley",
        {
            1: "六路径已执行（α₁/β₁ 见10.4；α₂/β₂/γ₁/γ₂ 见10.4b；shapley_allowed=False）",
            2: "效应均非零；α₁/β₁ 为 ΔD/yr、跨设备为百分比贡献，单位未统一，不能断言<3%可忽略；仅展示 D=0.5",
        },
    )
    replace_cell_text(
        doc, "分析项",
        "L4控制器",
        {
            1: "否定性结果+消融完成（l4_allowed=False）",
            2: "电池主导phi_C +5.3%（非-5.3%）；R1/R2消融实验已完成",
        },
    )

    print("== Table25 R2/R3 hashes ==")
    replace_cell_text(
        doc, "字段",
        "配置哈希",
        {
            1: GOV_CONFIG_HASH + "（治理口径，与统一结果台账 Table0 一致）",
            2: GOV_CONFIG_HASH + "（同左，共享 rectification_v2 input_contract）",
        },
    )
    replace_cell_text(
        doc, "字段",
        "代码哈希",
        {
            1: GOV_SCRIPT_HASH + "（治理口径，与统一结果台账 Table0 一致）",
            2: GOV_SCRIPT_HASH + "（同左，共享 rectification_v2 input_contract）",
        },
    )
    replace_cell_text(
        doc, "字段",
        "输入哈希",
        {1: "86e12930b6b08cab63dd00e12d5e00ab5afd0a3d1cc6e1122c90b1b52d207dd3（治理口径，与统一结果台账 Table0 一致）",
         2: "同左"},
    )

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
