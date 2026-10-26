# -*- coding: utf-8 -*-
"""生成MVP最小可行原型拆解文档"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
style.font.name = '宋体'; style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
for hname, hsize in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
    hs = doc.styles[hname]; hs.font.name = '黑体'; hs.font.size = Pt(hsize)
    hs.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    hs._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

def add_para(text, bold=False, color=None, size=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if color: r.font.color.rgb = RGBColor(*color)
    if size: r.font.size = Pt(size)
    return p

def add_note(text):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50); r.font.size = Pt(9.5)

def add_verify(text):
    p = doc.add_paragraph()
    r = p.add_run('【需人工核验】' + text)
    r.italic = True; r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B); r.font.size = Pt(10)

def add_table(headers, rows, col_widths=None):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Light Grid Accent 1'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, h in enumerate(headers):
        cell = t.rows[0].cells[ci]; cell.text = h
        for p in cell.paragraphs:
            for r in p.runs: r.bold = True; r.font.size = Pt(9.5)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri+1].cells[ci]; cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs: r.font.size = Pt(9)
    return t

# 标题
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('最小可行原型(MVP)拆解'); r.bold = True
r.font.size = Pt(20); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('功能模块 ｜ 数据准备 ｜ 测试流程 ｜ 验收指标 ｜ 文档清单')
r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x6A, 0x7A, 0x8A)
add_para('')
add_note('依据：05_阶段产出物/01_W1实施计划.docx（7阶段P1-P7 + 9核心模块 + 15步流程 + 16单元测试）')
add_note('原型目标：验证"配对反事实辨识+Shapley路径效应分解"方法在最小数据规模下可端到端运行并产出可验证的路径效应')
add_note('整理日期：2026-07-31')

doc.add_page_break()

# ============ MVP范围定义 ============
doc.add_heading('MVP范围定义', level=1)
add_table(
    ['维度', 'MVP范围', 'W1完整范围', '缩减理由'],
    [
        ['配置', 'C1(平衡) × 1 EMS(battery_first) = 1配置', '4容量×3EMS=12配置', 'C1最具代表性;battery_first逻辑最简'],
        ['窗口', 'W1(低PV) × 7天 = 1窗口', 'W1-W4 × 7天 = 4窗口', '低PV场景退化耦合压力最大'],
        ['健康运行', 'BOL + BAT D=0.50 + BAT D=0.90 = 3运行', 'BOL+3源×2健康=7运行', 'BAT源路径(α₁/β₁)数据支撑最强(A级)'],
        ['路径', 'α₁(BAT→FC) + β₁(BAT→EL) = 2路径', '6路径(α₁α₂β₁β₂γ₁γ₂)', 'BAT退化→FC/EL是优先路径'],
        ['路径效应分解', 'α₁路径×16联盟(完整Shapley); β₁仅AE不分解', 'N_eff×6锚点×16联盟', '验证联盟状态实现+Shapley闭合即可'],
        ['不确定性', '不执行(点估计)', '100组LHS+Bootstrap', 'MVP验证方法可行性,不确定性为P6任务'],
        ['年度验证', '不执行', '2017/2018×6配置×7健康=84次', 'MVP不涉及跨年度'],
        ['总运行数', '3次(观测) + 3次(反事实) + 16联盟 = 22次7天仿真', '~3192次', 'MVP约完整方案的0.7%'],
    ]
)
add_note('依据：产出物三§3.1计算量分解 + §3.2 V2锚点定义 + 产出物一§1.2六候选路径')
add_note('用途：在最小计算量下验证端到端方法链路完整可运行')
add_verify('MVP通过后,扩展到全部6路径仅需增加Run 4-7(FC/EL源),代码架构不变。')

# ============ 功能模块 ============
doc.add_heading('功能模块拆解（7个最小模块）', level=1)
add_table(
    ['模块', '功能', '输入', '输出', '依赖', '关键函数', '依据', '人工核验'],
    [
        ['M1 data_lite', '加载BAT退化数据+UCSD气象/负荷; 哈希校验',
         '280Ah aging.zip; UCSD 76CSV; frozen_parameters_v2.2.json',
         'data_registry_lite.json(哈希+文件路径)',
         '—', 'load_data(); compute_hash()',
         'W1模块1 data_registry', '280Ah 16时点数据完整性; UCSD W1窗口7天数据无缺失'],
        ['M2 hp_bat_lite', 'BAT E_usable(D)分段线性+R(D)幂律拟合; FC/EL用BOL常数',
         '280Ah 16时点容量+HPPC; IEEE PHM D=0极化曲线; NAICNO BOL面',
         'hp_bat.json(E_usable/R/P_max参数); hp_fc_bol.json; hp_el_bol.json',
         'M1', 'fit_E_usable(); fit_R(); get_fc_bol(); get_el_bol()',
         'W1模块4 health_performance', 'BAT方向单调≥98%; R(D)的a>0,b>0; FC/EL BOL参数来源标注'],
        ['M3 closed_loop_lite', 'PV+BAT+FC+EL+H₂ 15min时步闭环运行; 守恒检查',
         'hp_*.json; UCSD W1 7天数据; scenario_config(C1+battery_first)',
         'trajectory.parquet(15min功率/SOC/H₂/启停/爬坡)',
         'M2', 'run_pv(); run_bat(); run_fc(); run_el(); run_h2(); check_conservation()',
         'W1模块2 component_models', '功率残差≤10⁻⁶; 氢残差≤10⁻⁶; SOC/H₂在范围内'],
        ['M4 ems_bf_lite', 'battery_first EMS调度(伪代码实现)',
         'P_res=P_PV−P_load; SOC; H₂_level; P_max(D)',
         'ems_commands.json(每时步充放电/制氢/发电指令)',
         'M3', 'battery_first(); check_constraints()',
         'W1模块5 ems_rules(规则1)', '优先级链正确(充电→制氢→弃光; 放电→FC→shed); 爬坡≤Ramp_max'],
        ['M5 counterfactual_lite', '配对运行+源设备性能恢复+NC1+AE提取',
         'scenario_config; hp_*.json; ems_commands; 随机种子',
         'damage.csv(AE/TE); nc1_check.csv; pairing_hash.json',
         'M3/M4', 'run_pair(); restore_bat_perf(); extract_AE(); verify_hash(); run_nc1()',
         'W1模块6 counterfactual_engine', 'NC1残差≤10⁻⁴ EOL/年; 配对哈希一致; AE(α₁)>0且AE(β₁)>0'],
        ['M6 mediation_lite', 'α₁路径×16联盟Shapley值+闭合检查+交互指数',
         'damage.csv; counterfactual_lite; 联盟状态ON/OFF定义',
         'v2_mediation_lite.csv(16联盟值+φ_m+I_mn); shapley_closure.csv',
         'M5', 'compute_coalition_value(); shapley_value(); check_closure()',
         'W1模块7 mediation', 'Shapley闭合≤5%; 主导归因组物理可解释; OFF实现正确(g_i(D)→g_i(0))'],
        ['M7 report_lite', '结果CSV汇总+基本图表(路径AE/Shapley柱状图)',
         'damage.csv; v2_mediation_lite.csv',
         'mvp_report.csv; figures/ae_bar.png; figures/shapley_bar.png',
         'M5/M6', 'summarize_results(); plot_ae(); plot_shapley()',
         'W1工具D figure_generator', '图表标注MVP范围(C1/W1/battery_first/2路径)'],
    ],
    col_widths=[0.8, 1.5, 1.5, 1.3, 0.5, 1.2, 0.8, 1.2]
)
add_note('模块依赖链：M1→M2→M3→M4→M5→M6→M7,严格顺序不可跳层。M1/M2可部分并行(M2依赖M1数据但拟合可预跑)。')

# ============ 数据准备 ============
doc.add_heading('数据准备', level=1)
add_table(
    ['数据', '用途', '来源', '格式', '准备步骤', '人工核验'],
    [
        ['280Ah LFP容量数据', 'BAT E_usable(D)分段线性拟合', '数据文献/ Zenodo DOI:10.5281/zenodo.20813753', 'CSV(16行:cycle/capacity/Ohmic_res)',
         '解压aging.zip→提取LIC_aging_analysed.csv→确认16时点(cycle 0-700)', '时点数=16; capacity单调递减; Ohmic_res单调递增'],
        ['280Ah HPPC数据', 'BAT R(D)幂律拟合 R₀(1+aD^b)', '同上zip内HPPC文件', 'CSV(5温度×10SoC)',
         '提取35°C行→拟合R(D)=R₀(1+a×D^b)', 'a>0, b>0; NRMSE≤10%'],
        ['UCSD TradeStreet 15min', 'W1窗口7天PV/负荷数据', '02_data_processed/tradestreet_pipeline/76CSV', 'CSV(15min timestep)',
         '选取低PV周(如2017年1月)→截取7天→确认无缺失', '672行(7天×96步/天); PV/load列无NaN'],
        ['BLAST-Lite Sony/Murata', 'BAT容量交叉验证(风险1缓解)', 'NREL BLAST-Lite Python库', 'Python模型',
         'import sony_3ah→运行1yr@35°C/50%SOC→对比280Ah趋势', 'Sony容量方向与280Ah一致(均递减)'],
        ['IEEE PHM FC1/FC2 D=0', 'FC BOL极化曲线(目标设备性能参数)', '数据文献/ DOI:10.25666/DATAUBFC-2021-07-19', 'CSV(极化曲线)',
         '提取D=0时点→按j(A/cm²)归一化→拟合V=E_oc−Aln(j)−Bj−R_hfr·j', 'BOL极化方程拟合R²>0.95'],
        ['NAICNO BOL面(或替代)', 'EL BOL性能面(目标设备性能参数)', 'GitHub: NAICNO/wp7-UC2 或手动构建', 'Python模型/CSV',
         '验证仓库可访问→提取BOL V(j,T,p)面; 不可访问→用文献EIS数据手动构建', 'BOL面覆盖运行电流密度范围'],
        ['frozen_parameters_v2.2.json', 'C1配置参数(PV120/BAT135/FC37.5/EL35/H₂25)', '05_config/', 'JSON',
         '提取C1容量参数+SOC/H₂范围+Ramp_max+随机种子', '参数与产出物四§4.5一致'],
    ],
    col_widths=[1.3, 1.3, 1.5, 1.0, 1.8, 1.5]
)

# ============ 测试流程 ============
doc.add_heading('测试流程', level=1)
add_table(
    ['步骤', '测试名称', '模块', '操作', '预期', '通过标准'],
    [
        ['T1', '数据完整性', 'M1', '加载280Ah+UCSD数据→哈希校验', '16时点+672步无缺失', '哈希匹配+无NaN'],
        ['T2', 'BAT拟合', 'M2', '拟合E_usable(D)+R(D)', '容量非增+阻抗非减', '≥98%单调+NRMSE≤10%'],
        ['T3', 'BOL闭环', 'M3+M4', '7天BOL场景15min运行', '功率/氢守恒', '残差≤10⁻⁶+无越界'],
        ['T4', 'NC1预测试', 'M5', 'g_i(D_i)=g_i(0)强制BOL→测残差', 'ΔD_j≈0', '≤10⁻⁴ EOL/年'],
        ['T5', '配对一致性', 'M5', 'Run1(BOL) vs Run2(BAT D=0.50)字段哈希', '完全一致', '哈希相同'],
        ['T6', 'AE方向', 'M5', '提取AE(α₁)=ΔD_FC,Run2−ΔD_FC,Run1', 'AE(α₁)>0', 'BAT退化→FC损伤增加'],
        ['T7', 'AE方向', 'M5', '提取AE(β₁)=ΔD_EL,Run2−ΔD_EL,Run1', 'AE(β₁)>0', 'BAT退化→EL损伤增加'],
        ['T8', '联盟OFF', 'M6', 'OFF归因组g_i(D)→g_i(0)', '性能恢复BOL', 'E_usable/R/P_max=BOld值'],
        ['T9', 'Shapley闭合', 'M6', 'φ_1+φ_2+φ_3+φ_4 vs AE(α₁)', '之和=AE', '闭合≤5%'],
        ['T10', '物理可解释', 'M6', '主导归因组φ_max→物理检查', '与路径机制一致', '容量主导→缓冲减少→FC启停增加'],
        ['T11', '端到端', 'M1-M7', '全流程运行→报告生成', '22次仿真+CSV+图', '总时间<30分钟'],
    ],
    col_widths=[0.4, 1.0, 0.6, 2.0, 1.5, 1.5]
)
add_note('依据：W1产出物四 单元测试(16项)+集成测试(9项)的MVP子集')
add_note('用途：确保MVP每个模块通过测试后才集成到下一模块')
add_verify('T4(NC1)是方法核心的数值基础——若残差>10⁻⁴,需float128/Kahan后方可继续。T9(Shapley闭合)是方法核心的数学基础——若>5%,需查联盟值估计噪声。')

# ============ 验收指标 ============
doc.add_heading('验收指标', level=1)
add_table(
    ['指标', 'MVP通过标准', '依据', '测量方法'],
    [
        ['NC1残差', '≤10⁻⁴ EOL/年', '产出物四§4.6 验收门槛#10', 'g_i(D_i)=g_i(0)时ΔD_j的绝对值'],
        ['功率守恒', '≤额定功率×10⁻⁶', '同上 #1', 'ΣP_in − ΣP_out 的绝对值'],
        ['氢守恒', '≤累计氢吞吐×10⁻⁶', '同上 #2', 'ΣH₂_in − ΣH₂_out 的绝对值'],
        ['BAT方向单调', '≥98%时点满足 AND 连续违反≤3', '同上 #4', 'E_usable(D)非增; R(D)非减的时点比例'],
        ['BAT拟合NRMSE', '≤10%', '同上 #3', '拟合值vs实测值的归一化均方根误差'],
        ['AE方向正确', 'AE(α₁)>0 且 AE(β₁)>0', '产出物一§1.2 路径定义', 'BAT退化→FC/EL损伤增加(物理预期)'],
        ['Shapley闭合', '≤5%(解析值应≈0)', '产出物四§4.6 #11', '|φ_1+φ_2+φ_3+φ_4 − AE| / |AE|'],
        ['联盟OFF正确', 'OFF归因组性能=BOL值', '产出物二模块7 联盟状态规范', 'g_i(0)的输出=E_usable(0)/R(0)/P_max(0)'],
        ['物理可解释', '主导归因组与路径机制一致', '产出物六 风险4缓解', 'φ_max对应的物理含义检查'],
        ['端到端运行', '22次仿真总时间<30分钟', '产出物四§4.4 性能基准', 'wallclock时间测量'],
    ],
    col_widths=[1.2, 2.0, 1.8, 2.0]
)

# ============ 文档清单 ============
doc.add_heading('文档清单', level=1)
add_table(
    ['文档', '内容', '依据', '用途', '人工核验'],
    [
        ['MVP需求规格', '范围定义(1配置/1窗口/3运行/2路径/16联盟)+模块接口+依赖链', '本文档', '开发团队参照实现', '范围与W1产出物一致;缩减理由合理'],
        ['数据准备清单', '7项数据来源/格式/准备步骤', '本文档数据准备表', '数据加载前逐项确认', '280Ah 16时点完整;UCSD W1无缺失;NAICNO可访问或替代方案确定'],
        ['代码结构说明', '7模块文件结构+函数签名+import依赖', 'W1产出物二原型功能清单', '代码组织参考', '模块间接口(I/O文件名/格式)一致'],
        ['测试报告模板', '11项测试用例×结果×通过/未通过×备注', '本文档测试流程表', '每项测试填写结果', 'NC1和Shapley闭合两项必须通过'],
        ['结果验证表', 'AE(α₁)/AE(β₁)/φ_m/闭合误差/物理可解释性', '本文档验收指标', 'MVP结果与物理预期对比', 'AE方向正确+闭合≤5%+物理可解释'],
        ['MVP总结报告', '通过/未通过+问题清单+扩展建议', '本文档全部内容', '决策是否扩展到W1完整范围', '全部验收指标通过→扩展; 未通过→定位问题'],
    ],
    col_widths=[1.2, 2.0, 1.5, 1.5, 1.5]
)

doc.add_page_break()

# ============ 四方面评估 ============
doc.add_heading('四方面评估', level=1)

doc.add_heading('方面1：与阶段目标是否一致', level=2)
add_para('一致。MVP覆盖W1核心方法链(P1数据→P2健康-性能→P3闭环→P4反事实→P5归因),验证"配对反事实+Shapley路径效应分解"端到端可运行。'
         'MVP约完整方案的0.7%计算量(22次 vs 3192次),但方法链路完整。'
         '缩减项(LHS/年度/扩展/多配置多窗口)均为P6-P7任务,不影响核心方法验证。', size=10.5)

doc.add_heading('方面2：是否能支撑后续阶段', level=2)
add_para('基本支撑,但有3处扩展需注意：', bold=True, size=10.5)
for line in [
    '(1)MVP仅实现battery_first EMS→扩展到hydrogen_first/balanced需补2个EMS函数(M4扩展,约2天);',
    '(2)MVP仅BAT源设备→扩展到FC/EL源需补Run 4-7+对应健康-性能映射(M2扩展FC/EL从BOL常数到D→性能函数,约3-5天);',
    '(3)MVP无LHS/Bootstrap→扩展到V3不确定性需补lhs_sampler工具+100组采样(MVP代码架构支持,约2-3天);',
    '代码架构(模块化+JSON接口)可直接扩展,无需重构。',
]:
    add_para(line, size=10)

doc.add_heading('方面3：是否存在证据不足或概念不清', level=2)
add_para('存在2处需补充：', bold=True, size=10.5)
for line in [
    '(1)FC/EL作为目标设备在MVP中使用BOL常数(不做D→性能映射),但BOL极化曲线/性能面的来源需明确——'
    'IEEE PHM D=0时点数据和NAICNO BOL面(或替代)需在MVP启动前准备到位,否则闭环模型无法运行;',
    '(2)联盟状态OFF=g_i(0)的代码实现是MVP首次实现的关键逻辑——'
    '需在M6开发时编写专门的test_coalition_OFF单元测试,验证OFF归因组性能确实恢复为BOL值,'
    '而非仅修改参数变量名。',
]:
    add_para(line, size=10)

doc.add_heading('方面4：是否需要补充材料', level=2)
add_para('需补充3项：', bold=True, size=10.5)
for line in [
    '(1)FC BOL极化曲线拟合——从IEEE PHM FC1/FC2数据中提取D=0时点,按j归一化拟合极化方程,'
    '产出hp_fc_bol.json(V=E_oc−Aln(j)−Bj−R_hfr·j参数);',
    '(2)EL BOL性能面——验证NAICNO仓库可访问性;不可访问时用文献EIS数据手动构建V(j,T,p)面,'
    '产出hp_el_bol.json;',
    '(3)C1配置参数提取——从frozen_parameters_v2.2.json提取PV120/BAT135/FC37.5/EL35/H₂25及'
    'SOC/H₂范围/Ramp_max/随机种子,产出scenario_config_mvp.json。',
]:
    add_para(line, size=10)

# ============ 可执行修改任务 ============
doc.add_heading('可执行修改任务', level=1)

doc.add_heading('任务1：MVP启动前准备FC/EL BOL性能参数', level=3)
for line in [
    '(1) 从IEEE PHM FC1/FC2数据(DOI:10.25666/DATAUBFC-2021-07-19)提取D=0时点极化曲线;',
    '(2) 按电流密度j(A/cm²)归一化,拟合V=E_oc−A·ln(j)−B·j−R_hfr·j(R²>0.95);',
    '(3) 验证NAICNO/wp7-UC2仓库可访问性;不可访问则用PEMEL EIS文献数据手动构建V(j,T,p)BOL面;',
    '(4) 产出hp_fc_bol.json + hp_el_bol.json;',
    '(5) 预计0.5-1天。',
]:
    add_para(line, size=10)
add_verify('FC极化方程R²>0.95; EL BOL面覆盖运行电流密度范围[j_min, j_max]; NAICNO验证结果写入CP1报告。')

doc.add_heading('任务2：MVP启动前提取C1配置参数', level=3)
for line in [
    '(1) 读取05_config/frozen_parameters_v2.2.json;',
    '(2) 提取C1容量(PV120/BAT135/FC37.5/EL35/H₂25);',
    '(3) 提取SOC范围[SOC_min,SOC_max]、H₂范围[H₂_min,H₂_max]、Ramp_max(BAT/FC/EL);',
    '(4) 设定随机种子(如seed=42);',
    '(5) 产出scenario_config_mvp.json;',
    '(6) 预计0.5小时。',
]:
    add_para(line, size=10)
add_verify('参数与产出物四§4.5关键参数表一致; 随机种子固定且记录。')

doc.add_heading('任务3：M6归因模块开发时增加联盟OFF专项测试', level=3)
for line in [
    '(1) 在M6 mediation_lite中增加test_coalition_OFF单元测试;',
    '(2) 测试内容：对每个归因组(容量/效率/功率边界/库存控制),设为OFF→验证g_i输出=BOL值;',
    '(3) 验证方式：E_usable(D=0.50,OFF容量)=E_usable(D=0); R(D=0.50,OFF效率)=R(D=0); 等;',
    '(4) 若OFF输出≠BOL值→查代码中参数替换逻辑;',
    '(5) 预计0.5天(含调试)。',
]:
    add_para(line, size=10)
add_verify('4归因组×ON/OFF = 8种状态逐一验证; OFF输出与BOL值数值精确匹配(差<10⁻¹⁵)。')

doc.add_heading('任务4：MVP通过后编写扩展评估报告', level=3)
for line in [
    '(1) MVP全部11项测试通过后,编写扩展评估报告;',
    '(2) 评估内容：代码架构是否支持直接扩展(无重构); 扩展到6路径需增加的模块/函数;',
    '(3) 扩展工作量估计：+2 EMS(2天) + FC/EL D→性能(3-5天) + LHS工具(2-3天) = 7-10天;',
    '(4) 决策：MVP通过→启动W1完整执行; MVP未通过→定位问题→修复→重测;',
    '(5) 预计0.5天。',
]:
    add_para(line, size=10)
add_verify('扩展评估报告含代码架构审查+工作量估计+风险清单; 由开发负责人签字确认。')

# 保存
output_path = r'D:\matlab\C\09_docs\05_阶段产出物/02_MVP原型拆解.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
