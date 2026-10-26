# -*- coding: utf-8 -*-
"""Round-9 summary report fix: remove all stale wording found by full-text scan."""
from __future__ import annotations

from pathlib import Path

import docx

DOC = Path(r"d:\matlab\C\09_docs\工作区运行结果汇总.docx")


def set_para(doc, needle, new_text):
    for para in doc.paragraphs:
        if needle in para.text:
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_text
            print(f"[PARA] {needle[:24]!r} -> replaced")
            return True
    print(f"[WARN] para not found: {needle!r}")
    return False


def main() -> int:
    doc = docx.Document(str(DOC))

    # P8
    set_para(doc, "尚未断言：单位未统一之处",
             "口径说明：六路径单位已统一（effect_pct，D∈{0.3,0.5,0.7} 三水平），统一口径下 α₁/β₁ 量级大于跨设备路径"
             "（α₁ 最大 +24.33%）；成本参数边界仍缺，相关数值不作工程经济评估依据。")

    # P23 (two occurrences of 127/136)
    for para in doc.paragraphs:
        if "127 函数经 @parametrize 展开为 136 用例" in para.text:
            new_t = para.text.replace("127 函数经 @parametrize 展开为 136 用例", "140 函数经 @parametrize 展开为 149 用例")
            for run in para.runs:
                run.text = ""
            para.runs[0].text = new_t
            print("[PARA] P23 127/136 -> 140/149")

    # P41
    set_para(doc, "全部测试位于 08_tests/，由 pytest 驱动，地面真值 136 项 PASS",
             "全部测试位于 08_tests/，由 pytest 驱动，地面真值 149 项 PASS。按来源分解（与报告口径一致，经 pytest --collect-only 逐文件核验）：")

    # P50
    set_para(doc, "BAT 循环损伤仍用 BATKCYC_H 运行小时代理，未改 EFC/能量吞吐模型",
             "BAT 循环损伤已升级为 EFC/能量吞吐模型（closed_loop_lite.py 默认 BAT_CYCLE_MODEL='efc'，280Ah 标定"
             " a_d=5.0714e-4 D/EFC），不再使用运行小时代理。")

    # P52
    set_para(doc, "六路径效应单位未统一（ΔD/yr vs 百分比），「α₁/β₁ 显著大于跨设备路径」暂不成立",
             "六路径效应单位已统一（effect_pct 百分比贡献），统一口径下「α₁/β₁ 量级大于跨设备路径」成立"
             "（α₁ 最大 +24.33% @ hydrogen_dominant D=0.7；跨设备 |effect_pct|≤~3%）。")

    # P53
    set_para(doc, "嵌套双层传播（外层 Sobol × 内层季节移动块 Bootstrap）尚未启动；双层方差分解仅完成方法层",
             "嵌套双层传播（外层 Sobol × 内层季节移动块 Bootstrap）已完成（256×4000，闭合残差≤1.6e-16；"
             "内层为 2017 单年气象窗口，属“单年气象窗口不确定性”，不作三年气候稳健性外推）。")

    # P55 (next steps)
    set_para(doc, "① BAT 循环损伤改 EFC/能量吞吐模型 → 重跑 EOL 寿命与 NPC/LCOE",
             "执行顺序已完成（2026-08-04）：① BAT 循环损伤改 EFC/能量吞吐模型（已完成）→ 重跑 EOL 寿命与 NPC/LCOE（已完成）；"
             "② 统一六路径效应单位及三退化水平（D∈{0.3,0.5,0.7}）（已完成）；"
             "③ 嵌套双层传播 256×4000（已完成，单年气象窗口）；④ 成本敏感性与论文图表（F1–F5/T1–T5）定稿（已完成）。")

    # P49 limitation heading
    set_para(doc, "8.2　局限（保留，不作工程/物理论文结论依据）",
             "8.2　局限（成本参数边界等保留；已解决项见 6 节与 8.1 节）")

    # T1R7 (P1 table) 127/136
    for tb in doc.tables:
        if len(tb.rows) > 7 and "run_annual" in tb.rows[1].cells[1].text:
            tb.rows[7].cells[1].text = tb.rows[7].cells[1].text.replace(
                "127 函数经 @parametrize 展开为 136 用例", "140 函数经 @parametrize 展开为 149 用例")
            print("[CELL] T1R7 updated")
            break

    # T9R2 params -> EFC
    for tb in doc.tables:
        if len(tb.rows) > 2 and tb.rows[1].cells[0].text.strip() == "脚本" and \
           "run_degradation_uncertainty" in tb.rows[1].cells[1].text:
            for row in tb.rows:
                if row.cells[0].text.strip() == "输入数据":
                    row.cells[1].text = ("5 配置；6 退化参数（BAT_EFC_A_D/BAT_EFC_K/FC_K_RUN_H/FC_K_START/"
                                         "EL_K_RUN_H/EL_K_START）拉丁超立方 128 点")
                    print("[CELL] T9R2 params -> EFC")
            break

    # T14R7 pytest count
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                if "pytest，136 passed" in c.text:
                    c.text = c.text.replace("pytest，136 passed", "pytest，149 passed")
                    print("[CELL] T14 pytest -> 149")
    # fallback: "136 passed" anywhere
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                if "136 passed" in c.text:
                    c.text = c.text.replace("136 passed", "149 passed")
                    print("[CELL] 136 passed -> 149 passed")

    doc.save(str(DOC))
    print("saved:", DOC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
